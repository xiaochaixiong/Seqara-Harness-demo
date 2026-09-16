"""Final deep audit: edited real workbooks and interrupted output publication."""
import json
import unittest
from unittest.mock import patch
from pathlib import Path
import openpyxl
from docx import Document
import test_round2_repairs as fixtures
from nv_data_safety import staged_outputs, parse_headcount
import seqara_shell


class FinalDeepAudit(unittest.TestCase):
    setUp = fixtures.Round2Repairs.setUp
    book = fixtures.Round2Repairs.book
    frame = fixtures.Round2Repairs.frame
    read = fixtures.Round2Repairs.read
    invoice = fixtures.Round2Repairs.invoice
    handler = fixtures.Round2Repairs.handler

    def briefing(self, **options):
        return fixtures.ReportGenerator(self.messages.append, row_delay_sec=0, **options)

    def test_f01_user_notes_survive_three_invoice_runs(self):
        p = self.book('发票.xlsx', {'数据': [['发票项目','金额'],['住宿费',100]],
            '报销说明': [['人工核对记录'],['原件在财务室']]})
        for _ in range(3):
            p = self.invoice(p)
            self.assertEqual(self.read(p)['报销说明'][1][0], '原件在财务室')

    def test_f02_existing_metadata_named_user_sheet_is_preserved(self):
        for value in (None, '请保留我的备注'):
            p = self.book('同名工作表.xlsx', {'数据': [['发票项目','金额'],['住宿费',100]],
                '_Seqara处理信息': [['用户的汇总说明'],[],['备注',value]]})
            for _ in range(2):
                p = self.invoice(p)
                self.assertEqual(self.read(p)['_Seqara处理信息'][0][0], '用户的汇总说明')

    def test_f03_correction_does_not_leak_across_batches(self):
        p = self.book('两批.xlsx', {title: [['序号','发票项目','金额'],[1,item,100]]
            for title,item in [('批次甲','住宿费'),('批次乙','火车票')]})
        p = self.invoice(p)
        w = openpyxl.load_workbook(p)
        for ws in w:
            h = fixtures.invoices.build_header_map(ws,fixtures.invoices.detect_header_row(ws))
            if '预测分类标签' in h:
                for row in range(2,ws.max_row+1):
                    if ws.cell(row,h['发票项目']).value == '住宿费':
                        ws.cell(row,h['分类标签']).value = '服务费'
        w.save(p); w.close()
        p = self.invoice(p)
        rows = self.read(p)['分类&分组']; h = rows[0]
        labels = {r[h.index('发票项目')]:r[h.index('分类标签')] for r in rows[1:] if r[h.index('发票项目')]}
        self.assertEqual(labels['住宿费'], '服务费')
        self.assertEqual(labels['火车票'], '交通费')

    def test_f04_total_word_in_notes_does_not_drop_invoice(self):
        p = self.book('备注.xlsx', {'数据': [['发票项目','金额','备注'],['住宿费',100,'合计'],['会议费',200,None],['总计',300,None]]})
        self.assertEqual(self.read(self.invoice(p))['分组统计'][1][2],300)

    def test_f05_malformed_amount_separators_fail_at_source_cell(self):
        for value in ('1,2','12,34','1,,200','1 2'):
            p = self.book('金额.xlsx', {'明细': [['发票项目','金额'],['住宿费',value]]})
            with self.assertRaisesRegex(ValueError, 'B2'):
                self.invoice(p)
        for value,expected in [('￥1,200.50元',1200.5),('(1,200.50)',-1200.5),(' -12.50 ',-12.5)]:
            self.assertEqual(fixtures.invoices.parse_amount(value),expected)

    def test_f06_internal_spaces_cannot_join_two_counts(self):
        for value in ('1 2','12\t34','1\n20'):
            with self.assertRaises(ValueError): parse_headcount(value)
        self.assertEqual(parse_headcount(' 1,200 '),1200)

    def test_f07_briefing_count_matches_statistics(self):
        p = self.frame('简报.xlsx',[{**fixtures.BASE,'活动内容':'本地正文','服务教师（人数）':'1,200'}])
        result = self.briefing(keep_activity_meta=True).generate(p,self.home/'brief')
        text = '\n'.join(x.text for x in Document(result['artifacts'][0]).paragraphs)
        self.assertIn('服务教师：1200人',text)
        p = self.frame('错误简报.xlsx',[{**fixtures.BASE,'活动内容':'本地正文','服务教师（人数）':'1.5'}])
        result = self.briefing(keep_activity_meta=True).generate(p,self.home/'bad')
        self.assertEqual(result['success'],0)
        self.assertTrue(result['failed'])
        self.assertEqual(result['artifacts'],[])

    def test_f08_merged_briefing_save_failure_reports_reason(self):
        from docx.document import Document as DocClass
        p = self.frame('简报.xlsx',[{**fixtures.BASE,'活动内容':'本地正文'}])
        with patch.object(DocClass,'save',side_effect=PermissionError('文件被占用')):
            result = self.briefing().generate(p,self.home/'brief')
        self.assertTrue(result['failed'],result)
        self.assertIn('文件被占用',str(result['failed']))
        self.assertEqual(list((self.home/'brief').iterdir()),[])

    def test_f09_uppercase_csv_extension(self):
        import pandas as pd
        p = self.home/'简报.CSV'
        pd.DataFrame([{**fixtures.BASE,'活动内容':'本地正文'}]).to_csv(p,index=False,encoding='utf-8-sig')
        result = self.briefing().generate(p,self.home/'brief')
        self.assertEqual(result['success'],1,result)

    def test_f10_cancellation_during_last_activity_prevents_publish(self):
        p = self.frame('简报.xlsx',[{**fixtures.BASE,'活动内容':'本地正文'}])
        for split in (False,True):
            gen = self.briefing(split_by_activity=split)
            append = gen._append_activity_section
            def stop(*args,**kwargs):
                ok = append(*args,**kwargs);gen.stop_flag=True;return ok
            with patch.object(gen,'_append_activity_section',side_effect=stop):
                result = gen.generate(p,self.home/str(split))
            self.assertEqual(result['status'],'cancelled')
            self.assertEqual(result['artifacts'],[],result)

    def test_f11_missing_staging_file_leaves_no_empty_deliverable(self):
        a,b = self.home/'A.xlsx',self.home/'B.csv'
        with self.assertRaises(FileNotFoundError):
            with staged_outputs(a,b) as paths:
                paths[0].write_bytes(b'complete workbook')
        self.assertFalse(a.exists())
        self.assertFalse(b.exists())

    def test_f12_invalid_sidebar_cache_recovers(self):
        for payload in ([],42,{'version':1,'html':'<aside/>','width':'损坏'},
            {'version':1,'html':'<aside/>','width':None},
            {'version':1,'html':'<aside/>','width':float('inf')},
            {'version':1,'html':'<aside/>','width':280,'css':{}}):
            (self.home/'sidebar-cache.json').write_text(json.dumps(payload),encoding='utf8')
            cached = seqara_shell.load_cache(self.home)
            if cached is not None:
                self.assertIsInstance(cached.get('css',''),str)
                self.assertTrue(52 <= int(cached.get('width',280)) <= 400)

    def test_f13_replaced_print_file_is_a_new_job(self):
        import os
        p = self.home/'下载.pdf'; p.write_bytes(b'first document')
        h = self.handler()
        with patch.object(fixtures.printing,'wait_file_size_stable',return_value=True), patch.object(fixtures.printing,'dispatch_print',return_value=True) as printer:
            h._try_print(str(p.resolve()))
            h._try_print(str(p.resolve()))
            self.assertEqual(printer.call_count,1)
            p.write_bytes(b'a new document with the same download filename')
            h._try_print(str(p.resolve()))
            self.assertEqual(printer.call_count,2)


if __name__ == '__main__': unittest.main(verbosity=2)
