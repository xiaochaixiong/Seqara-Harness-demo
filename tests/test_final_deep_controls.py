"""Compatibility and retry acceptance for the final audit fixes."""
import json
import unittest
from pathlib import Path
from unittest.mock import patch
import openpyxl
import test_final_deep_audit as deep
import test_user_audit_repairs as old
import test_round2_repairs as fixtures
from nv_data_safety import staged_outputs


class FinalControls(unittest.TestCase):
    setUp = fixtures.Round2Repairs.setUp
    book = fixtures.Round2Repairs.book
    frame = fixtures.Round2Repairs.frame
    read = fixtures.Round2Repairs.read
    invoice = fixtures.Round2Repairs.invoice
    briefing = deep.FinalDeepAudit.briefing

    def test_atomic_batch_collision_rollback_and_retry(self):
        a,b = self.home/'A.xlsx',self.home/'B.csv'
        with self.assertRaises(FileExistsError):
            with staged_outputs(a,b) as paths:
                paths[0].write_bytes(b'A');paths[1].write_bytes(b'B')
                b.write_bytes(b'user file created during processing')
                self.assertFalse(a.exists())
        self.assertFalse(a.exists())
        self.assertEqual(b.read_bytes(),b'user file created during processing')
        with staged_outputs(a) as (staged,):staged.write_bytes(b'complete retry')
        self.assertEqual(a.read_bytes(),b'complete retry')

    def test_distinct_targets_with_same_basename(self):
        (self.home/'another').mkdir()
        a,b=self.home/'A.xlsx',self.home/'another/A.xlsx'
        with staged_outputs(a,b) as paths:
            paths[0].write_bytes(b'A');paths[1].write_bytes(b'B')
        self.assertEqual((a.read_bytes(),b.read_bytes()),(b'A',b'B'))

    def legacy_workbook(self,duplicate=False):
        rows=[['序号','发票项目','金额'],[1,'住宿费',100]]
        if duplicate:rows.append([1,'火车票',100])
        p=self.invoice(self.book('旧版.xlsx',{'数据':rows,'用户备注':[['请保留这页']]}))
        w=openpyxl.load_workbook(p)
        for ws in w:
            h=fixtures.invoices.build_header_map(ws,fixtures.invoices.detect_header_row(ws))
            if fixtures.invoices.INVOICE_ROW_ID in h:ws.delete_cols(h[fixtures.invoices.INVOICE_ROW_ID])
        m=w['_Seqara处理信息'];m['A1']='Seqara invoice sources v1'
        m['B2']=json.dumps([ws.title for ws in w if ws.title not in ('数据',m.title)],ensure_ascii=False)
        m['B3']=json.dumps({'seq:1|amt:100.0':'服务费'},ensure_ascii=False)
        w.save(p);w.close();return p

    def test_upgrade_legacy_metadata_keeps_notes_and_correction(self):
        for note_name in ('用户备注','分组统计','未分类_待处理'):
            p=self.legacy_workbook()
            if note_name!='用户备注':
                w=openpyxl.load_workbook(p)
                w[note_name].title=note_name+'_旧生成'
                w['用户备注'].title=note_name
                m=w['_Seqara处理信息']
                m['B2']=json.dumps([ws.title for ws in w if ws.title not in ('数据',m.title)],ensure_ascii=False)
                w.save(p);w.close()
            for _ in range(2):
                p=self.invoice(p);sheets=self.read(p)
                self.assertEqual(sheets[note_name][0][0],'请保留这页')
                h=sheets['分类&分组'][0];self.assertEqual(sheets['分类&分组'][1][h.index('分类标签')],'服务费')

    def test_ambiguous_legacy_correction_requires_review(self):
        p=self.legacy_workbook(True);before=p.read_bytes()
        with self.assertRaisesRegex(ValueError,'无法确定人工分类'):self.invoice(p)
        self.assertEqual(p.read_bytes(),before)

    def test_identity_hidden_and_stable_after_source_reorder(self):
        p=self.invoice(self.book('排序.xlsx',{'原始':[['序号','发票项目','金额'],[1,'住宿费',100],[2,'火车票',200]]}))
        w=openpyxl.load_workbook(p)
        for ws in w:
            h=fixtures.invoices.build_header_map(ws,fixtures.invoices.detect_header_row(ws))
            if fixtures.invoices.INVOICE_ROW_ID in h:
                self.assertTrue(ws.column_dimensions[openpyxl.utils.get_column_letter(h[fixtures.invoices.INVOICE_ROW_ID])].hidden)
            if '预测分类标签' in h:
                for row in range(2,ws.max_row+1):
                    if ws.cell(row,h['发票项目']).value=='住宿费':ws.cell(row,h['分类标签']).value='服务费'
        ws=w['原始'];a,b=list(ws.values)[1:]
        for r,values in [(2,b),(3,a)]:
            for c,v in enumerate(values,1):ws.cell(r,c).value=v
        w.save(p);w.close()
        data=self.read(self.invoice(p))['分类&分组'];h=data[0]
        self.assertEqual(next(r[h.index('分类标签')] for r in data[1:] if r[h.index('发票项目')]=='住宿费'),'服务费')

    def test_cancel_keeps_only_previously_completed_activity(self):
        p=self.frame('取消.xlsx',[{**fixtures.BASE,'活动名称':x,'活动内容':'正文'} for x in ('完成活动','取消活动')])
        gen=self.briefing(split_by_activity=True);append=gen._append_activity_section
        def stop(doc,row,**kwargs):
            ok=append(doc,row,**kwargs)
            if kwargs['index']==1:gen.stop_flag=True
            return ok
        with patch.object(gen,'_append_activity_section',side_effect=stop):r=gen.generate(p,self.home/'out')
        self.assertEqual(r['status'],'cancelled');self.assertEqual(r['success'],1)
        self.assertEqual(len(r['artifacts']),1);self.assertIn('完成活动',r['artifacts'][0])

    def test_large_multisheet_invoice_totals_remain_exact_on_rerun(self):
        sheets={f'批次{i}':[['序号','发票项目','金额']]+[[j,'住宿费',j+0.25] for j in range(1,31)]+[['合计',None,472.5]] for i in range(10)}
        p=self.book('多批次.xlsx',sheets)
        for _ in range(2):
            p=self.invoice(p);rows=self.read(p)['分组统计']
            self.assertEqual(sum(r[2] for r in rows[1:]),4725)


class FinalWorkerControls(unittest.TestCase):
    setUp=fixtures.Round2Repairs.setUp
    book=fixtures.Round2Repairs.book
    frame=fixtures.Round2Repairs.frame
    read=fixtures.Round2Repairs.read
    invoke=old.UserAuditRepairs.invoke

    def test_release_preserves_notes_and_invoice_amounts(self):
        p=self.book('重复.xlsx',{'数据':[['发票项目','金额','备注'],['住宿费',100,'合计']],
                               '用户备注':[['必须保留']]})
        for _ in range(2):
            result=self.invoke('invoices',{'source':str(p),'company':'测试公司','project':'最终验收'})
            p=Path(next(x for x in result['artifacts'] if x.endswith('.xlsx')))
            data=self.read(p);self.assertIn('用户备注',data);self.assertEqual(data['分组统计'][1][2],100)

    def test_release_rejects_malformed_amount_and_counts(self):
        p=self.book('坏金额.xlsx',{'原票':[['发票项目','金额'],['住宿费','1,2']]})
        self.assertIn('B2',str(self.invoke('invoices',{'source':str(p),'company':'测试公司','project':'最终验收'},False)))
        p=self.frame('坏人数.xlsx',[{**fixtures.BASE,'服务教师（人数）':'1 2'}])
        self.invoke('statistics',{'source':str(p)},False)

    def test_release_briefing_comma_counts(self):
        from docx import Document
        p=self.frame('人数.xlsx',[{**fixtures.BASE,'服务教师（人数）':'1,200','活动内容':'本地正文'}])
        result=self.invoke('briefing',{'source':str(p)})
        text='\n'.join(p.text for p in Document(result['artifacts'][0]).paragraphs)
        self.assertIn('本地正文',text)
        self.assertEqual(result['briefing']['success'],1)
        p=self.frame('无效人数.xlsx',[{**fixtures.BASE,'服务教师（人数）':'1.5','活动内容':'本地正文'}])
        self.assertIn('非负整数',str(self.invoke('briefing',{'source':str(p)},False)))


if __name__=='__main__':unittest.main(verbosity=2)
