"""Regression scenarios from the 0.8 ordinary-user audit; no real network/printing."""
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

import openpyxl
import pandas as pd
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'native'))
from nv_briefing_core import ExcelSplitter, ReportGenerator
from nv_dashboard_core import build_dashboard_snapshot, DashboardStateStore, _parse_date
import nv_business_core as business


class UserAuditRepairs(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.home = Path(self.tmp.name)

    def book(self, name, sheets):
        target = self.home / name
        book = openpyxl.Workbook(); book.remove(book.active)
        for name, rows in sheets.items():
            ws = book.create_sheet(name)
            for row in rows: ws.append(row)
        book.save(target); book.close()
        return target

    def read(self, path):
        book = openpyxl.load_workbook(path)
        try: return {s.title: list(s.values) for s in book}
        finally: book.close()

    def invoke(self, name, args, ok=True):
        command = [os.environ['NV_FIX_WORKER_PATH']] if os.environ.get('NV_FIX_WORKER_PATH') else [sys.executable, str(ROOT/'backend/worker.py')]
        result = subprocess.run(command, input=json.dumps({'name': name, 'args': args}),
            capture_output=True, encoding='utf8', timeout=60,
            env={**os.environ, 'NV_TOOLKIT_DATA_DIR': str(self.home/'state'), 'SEQARA_OFFLINE': '1', 'PYTHONIOENCODING': 'utf-8'})
        events = [json.loads(line, parse_constant=lambda x: self.fail('Non-JSON number: '+x))
                  for line in result.stdout.splitlines() if line.startswith('{')]
        self.assertTrue(events, result.stderr)
        self.assertEqual(events[-1]['type'], 'result' if ok else 'error', events)
        self.assertEqual(result.returncode, 0 if ok else 1)
        return events[-1].get('value', events[-1])

    def test_split_keeps_source_and_existing_outputs(self):
        source = self.book('甲大学.xlsx', {'数据': [['学校名称','金额'],['甲大学',10],['乙大学',20]]})
        before = source.read_bytes()
        existing = self.home/'乙大学.xlsx'; existing.write_bytes(b'user-owned output')
        result = ExcelSplitter(lambda _: None).split(source, self.home)
        self.assertEqual(result['success'], 2)
        self.assertEqual(source.read_bytes(), before)
        self.assertEqual(existing.read_bytes(), b'user-owned output')
        self.assertEqual(self.read(self.home/'乙大学_2.xlsx')['数据'][1], ('乙大学',20))

    def test_split_all_sheets_and_excludes_unfiltered_sheet(self):
        source = self.book('总表.xlsx', {'第一期':[['学校名称'],['甲大学']], '第二期':[['学校名称'],['乙大学']], '内部汇总':[['院校'],['乙大学']]})
        result = self.invoke('split', {'source': str(source)})
        self.assertEqual(result['count'], 2)
        files = {Path(p).name: self.read(p) for p in result['artifacts']}
        self.assertNotIn('内部汇总', files['甲大学.xlsx'])
        self.assertEqual(files['乙大学.xlsx']['第二期'][1], ('乙大学',))
        self.assertIn('内部汇总', self.read(source))

    def test_split_remaps_relative_and_absolute_same_row_formulas(self):
        source = self.book('公式.xlsx', {'数据':[['学校名称','金额','加倍'],['甲大学',10,'=B2*2'],['乙大学',20,'=$B$3*2'],['甲大学',30,'=B4*2']]})
        result = self.invoke('split', {'source': str(source)})
        files = {Path(p).name: self.read(p)['数据'] for p in result['artifacts']}
        self.assertEqual(files['乙大学.xlsx'][1][2], '=$B$2*2')
        self.assertEqual(files['甲大学.xlsx'][2][2], '=B3*2')

    def test_split_rejects_uncalculated_cross_row_formula(self):
        source = self.book('公式.xlsx', {'数据':[['学校名称','金额','总额'],['甲大学',10,'=SUM(B2:B3)'],['乙大学',20,None]]})
        result = ExcelSplitter(lambda _: None).split(source, self.home/'out')
        self.assertTrue(any('重算' in error for error in result['failed']))
        self.assertFalse((self.home/'out/甲大学.xlsx').exists())

    def test_clean_does_not_fill_unrelated_activity(self):
        common = {'学院名称':'艺术学院','活动类型':'师资培训','申请人部门':'驻校办公室/甲大学'}
        a = {'活动名称':'培训甲','审批编号':'202609010001','开始时间':'2026-09-01','服务教师（人数）':10,'服务学生（人数）':20,'活动总额':1000, **common}
        b = {'活动名称':'竞赛乙','审批编号':'202609010002','开始时间':None,'服务教师（人数）':None,'服务学生（人数）':None,'活动总额':None, **common}
        source=self.home/'申请.xlsx'; pd.DataFrame([a,b]).to_excel(source,index=False)
        reim=self.home/'报销.xlsx'; pd.DataFrame([{'活动名称':'培训甲','关联申请单':'202609010001','新闻链接':'https://example.invalid/a'}, {'活动名称':'竞赛乙','关联申请单':'202609010002','新闻链接':None}]).to_excel(reim,index=False)
        result=self.invoke('clean', {'source':str(source),'reimbursement':str(reim)})
        output=next(p for p in result['artifacts'] if p.endswith('.xlsx'))
        rows=pd.read_excel(output).set_index('活动名称')
        for col in ('活动时间','服务教师（人数）','服务学生（人数）','费用','新闻链接'):
            self.assertTrue(pd.isna(rows.loc['竞赛乙',col]), (col,rows.loc['竞赛乙',col]))

    def test_clean_still_expands_real_merged_cells(self):
        source=self.book('合并.xlsx',{'数据':[['活动名称','人数'],['培训',10],[None,None]]})
        w=openpyxl.load_workbook(source);w.active.merge_cells('A2:A3');w.active['B3']=20;w.save(source);w.close()
        frame=business._read_activity_excel(source)
        self.assertEqual(frame.iloc[1,0],'培训')

    def invoice(self, amount):
        source=self.book('发票.xlsx', {'数据':[['发票项目','金额（元）','发票号码'],['*住宿服务*住宿费',amount,'TEST001']]})
        return {'source':str(source),'company':'测试公司','project':'测试项目','base':1000,'jitter':0}

    def test_invoice_currency_preserves_amount(self):
        result=self.invoke('invoices',self.invoice('￥1,200.00'))
        rows=self.read(result['artifacts'][0])['分组统计']
        self.assertEqual(rows[1][2],1200)

    def test_invoice_unknown_amount_is_error(self):
        result=self.invoke('invoices',self.invoice('待确认'),False)
        self.assertIn('B2',result['message'])

    def test_invoice_uncalculated_formula_is_error(self):
        result=self.invoke('invoices',self.invoice('=100+200'),False)
        self.assertIn('重算',result['message'])

    def dashboard(self, counts):
        source=self.home/'活动.xlsx'
        pd.DataFrame([{'学校名称':'甲大学','活动名称':name,'活动类型':'培训','活动时间':'2026-09-01','服务教师（人数）':count,'服务学生（人数）':0} for name,count in counts]).to_excel(source,index=False)
        return source

    def test_dashboard_reorder_keeps_exclusions_and_unexclude(self):
        source=self.dashboard([('甲',10),('乙',30)])
        db=self.home/'state.sqlite3'; store=DashboardStateStore(db)
        snapshot=build_dashboard_snapshot(source,None,db)
        first=next(a for a in snapshot.activities if a.activity_name=='甲')
        store.set_coverage_excluded(first.approval_id,'teacher',True)
        self.dashboard([('乙',30),('甲',10)])
        result=build_dashboard_snapshot(source,None,db)
        self.assertEqual(result.teacher_total,30)
        store.set_coverage_excluded(first.approval_id,'teacher',False)
        self.assertEqual(build_dashboard_snapshot(source,None,db).teacher_total,40)

    def test_dashboard_migrates_existing_exclusion(self):
        source=self.dashboard([('甲',10)]);db=self.home/'state.sqlite3';store=DashboardStateStore(db)
        identity='\x1f'.join(['甲大学','甲','培训','2026-09-01','1'])
        old='CLN-'+hashlib.sha1(identity.encode()).hexdigest()[:12].upper()
        store.set_coverage_excluded(old,'teacher',True)
        snapshot=build_dashboard_snapshot(source,None,db)
        self.assertEqual(snapshot.teacher_total,0)
        self.assertNotIn((old,'teacher'),store.coverage_exclusions())
        store.set_coverage_excluded(snapshot.activities[0].approval_id,'teacher',False)
        self.assertEqual(build_dashboard_snapshot(source,None,db).teacher_total,10)

    def test_dashboard_rejects_invalid_counts_with_quality_and_valid_json(self):
        source=self.dashboard([('无穷','inf'),('负数',-10),('小数',1.5),('千分位','1,200')])
        result=self.invoke('dashboard',{'source':str(source)})['dashboard']
        self.assertEqual(result['teachers'],1200)
        self.assertEqual(len(result['quality']),3)

    def test_compact_dates(self):
        self.assertEqual(str(_parse_date(20260901)),'2026-09-01')
        self.assertEqual(str(_parse_date('20260901')),'2026-09-01')
        self.assertIsNone(_parse_date('20260230'))

    def briefing(self, rows, separate=False):
        source=self.home/'简报.xlsx';pd.DataFrame(rows).to_excel(source,index=False)
        logs=[];gen=ReportGenerator(logs.append, row_delay_sec=0, split_by_activity=separate)
        with patch.dict(os.environ, {'SEQARA_OFFLINE':'1'}), patch.object(gen.session,'get',side_effect=AssertionError('offline network')):
            result=gen.generate(str(source),str(self.home/'docs'))
        texts=['\n'.join(p.text for p in Document(f).paragraphs) for f in result['artifacts']]
        return result,texts,logs

    def test_briefing_local_content_without_link(self):
        result,texts,_=self.briefing([{'学校名称':'甲大学','活动名称':'培训','活动内容':'本地活动正文'}])
        self.assertEqual(result['status'],'completed');self.assertIn('本地活动正文',texts[0])

    def test_briefing_all_failed_produces_no_success(self):
        result,texts,logs=self.briefing([{'学校名称':'甲大学','活动名称':'培训','新闻链接':'https://example.invalid'}])
        self.assertEqual(result['status'],'failed');self.assertFalse(texts)
        self.assertFalse(any('成功' in log for log in logs))

    def test_briefing_partial_failure_is_explicit(self):
        result,texts,logs=self.briefing([{'学校名称':'甲大学','活动名称':'培训','活动内容':'有效正文'}, {'学校名称':'甲大学','活动名称':'失败活动','新闻链接':'https://example.invalid'}])
        self.assertEqual(result['status'],'partial');self.assertEqual(result['success'],1)
        self.assertNotIn('失败活动',texts[0]);self.assertTrue(any('部分成功' in log for log in logs))

    def test_briefing_requires_school(self):
        result,texts,_=self.briefing([{'活动名称':'培训','活动内容':'正文'}])
        self.assertFalse(texts);self.assertIn('学校',result['failed'][0])

    def test_briefing_multi_school_retains_attribution(self):
        rows=[{'学校名称':school,'活动名称':school+'培训','活动内容':'正文'} for school in ('甲大学','乙大学')]
        result,texts,_=self.briefing(rows)
        self.assertTrue(texts[0].startswith('多院校'))
        self.assertIn('乙大学 · 乙大学培训',texts[0])
        result,texts,_=self.briefing(rows,True)
        self.assertTrue(any(t.startswith('乙大学示例') for t in texts))

    def test_worker_briefing_local_and_error(self):
        source=self.home/'简报.xlsx';pd.DataFrame([{'学校名称':'甲大学','活动名称':'培训','活动内容':'本地正文'}]).to_excel(source,index=False)
        result=self.invoke('briefing',{'source':str(source)})
        self.assertEqual(result['briefing']['status'],'completed')
        pd.DataFrame([{'学校名称':'甲大学','新闻链接':'https://example.invalid'}]).to_excel(source,index=False)
        self.assertIn('简报生成失败',self.invoke('briefing',{'source':str(source)},False)['message'])

    def test_empty_statistics_actionable(self):
        source=self.home/'空.xlsx';pd.DataFrame(columns=['学校名称','活动类型']).to_excel(source,index=False)
        self.assertIn('没有可统计的数据',self.invoke('statistics',{'source':str(source)},False)['message'])

    def test_native_backend_core_match(self):
        for name in ('nv_briefing_core.py','nv_business_core.py','nv_classify_core.py','nv_dashboard_core.py','nv_dashboard_widgets.py'):
            self.assertEqual((ROOT/'native'/name).read_bytes(),(ROOT/'backend'/name).read_bytes(),name)


if __name__ == '__main__': unittest.main(verbosity=2)
