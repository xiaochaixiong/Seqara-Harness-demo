"""Next-stage round 2: edited outputs, continuous workflows, export and retry state."""
import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch
import openpyxl
import pandas as pd
import test_round2_repairs as fixtures
from nv_dashboard_core import filter_dashboard_snapshot
from nv_dashboard_export import export_dashboard_view_xlsx


class NextRound2(unittest.TestCase):
    setUp=fixtures.Round2Repairs.setUp
    book=fixtures.Round2Repairs.book
    frame=fixtures.Round2Repairs.frame
    read=fixtures.Round2Repairs.read
    invoice=fixtures.Round2Repairs.invoice
    handler=fixtures.Round2Repairs.handler

    def test_n2_01_02_manual_invoice_correction_survives_rerun(self):
        for project in ('测试项目','很长的项目名称'*12):
            p=self.book('原票.xlsx',{'原始':[['发票项目','金额'],['ZXQ未知项目',100]]})
            p=self.invoice(p,project=project)
            w=openpyxl.load_workbook(p)
            for ws in w:
                headers=fixtures.invoices.build_header_map(ws,fixtures.invoices.detect_header_row(ws))
                if '预测分类标签' in headers:
                    ws.cell(3,headers['分类标签']).value='服务费'
            w.save(p);w.close()
            p=self.invoice(p,project=project)
            rows=self.read(p)['分类&分组']
            header=rows[0];self.assertEqual(rows[1][header.index('分类标签')],'服务费',project)

    def test_n2_03_new_sheet_added_to_processed_workbook(self):
        p=self.book('原票.xlsx',{'原始':[['发票项目','金额'],['住宿费',100]]})
        p=self.invoice(p);w=openpyxl.load_workbook(p);s=w.create_sheet('追加批次')
        s.append(['发票项目','金额']);s.append(['住宿费',200]);w.save(p);w.close()
        self.assertEqual(self.read(self.invoice(p))['分组统计'][1][2],300)

    def test_n2_04_dashboard_export_cannot_replace_source(self):
        p=self.frame('看板源.xlsx',[fixtures.BASE]);before=p.read_bytes()
        s=fixtures.build_dashboard_snapshot(p,None,self.home/'state.db')
        with self.assertRaisesRegex(ValueError,'输入|源文件'):
            export_dashboard_view_xlsx(s,filter_dashboard_snapshot(s),p)
        self.assertEqual(p.read_bytes(),before)

    def test_n2_05_stat_save_failure_leaves_no_partial_file(self):
        p=self.frame('统计源.xlsx',[fixtures.BASE])
        original=openpyxl.Workbook.save
        def fail_style(w,p):
            if w.active['A1'].font.bold:raise PermissionError('模拟样式保存失败')
            return original(w,p)
        with patch.object(openpyxl.Workbook,'save',new=fail_style):
            result=fixtures.biz.generate_project_stat_tables(p,self.home/'out')
        self.assertIsNone(result)
        self.assertEqual(list((self.home/'out').glob('*')),[])

    def test_n2_06_malformed_comma_headcount_rejected_consistently(self):
        from nv_data_safety import parse_headcount
        from nv_dashboard_core import _parse_number
        for value in ('1,2','12,34','1,,200'):
            with self.assertRaises(ValueError):parse_headcount(value)
            self.assertTrue(_parse_number(value)[1])
        self.assertEqual(parse_headcount('1,200'),1200)

    def test_n2_09_explicit_approval_id_survives_date_and_count_edits(self):
        p=self.frame('看板.xlsx',[{**fixtures.BASE,'审批编号':'202609010001'}]);db=self.home/'state.db'
        s=fixtures.build_dashboard_snapshot(p,None,db)
        fixtures.DashboardStateStore(db).set_coverage_excluded(s.activities[0].approval_id,'teacher',True)
        self.frame('看板.xlsx',[{**fixtures.BASE,'审批编号':'202609010001','活动时间':'2026-08-01','服务教师（人数）':30}])
        s=fixtures.build_dashboard_snapshot(p,None,db);self.assertEqual(s.teacher_total,0)

    def test_n2_10_zip_duplicate_destination_is_rejected(self):
        p=self.home/'重复.zip'
        with zipfile.ZipFile(p,'w') as z:z.writestr('A.pdf',b'A');z.writestr('a.PDF',b'B')
        files,error=fixtures.printing.safe_extract_zip(p,self.home/'unzip')
        self.assertTrue(error,'Windows 上同名条目会覆盖并重复打印')
        self.assertEqual(files,[])

    def test_n2_11_cancelled_invoice_creates_no_deliverable(self):
        import threading
        p=self.book('原票.xlsx',{'原始':[['发票项目','金额'],['住宿费',100]]})
        cancel=threading.Event();cancel.set()
        with self.assertRaisesRegex(RuntimeError,'取消'):self.invoice(p,cancel_flag=cancel)
        self.assertEqual(list((self.home/'invoice-out').glob('*.xlsx')),[])


class NextRound2Widgets(unittest.TestCase):
    setUpClass=fixtures.Round2WidgetRepairs.setUpClass

    def test_n2_07_missing_selected_month_keeps_explicit_filter(self):
        import tempfile
        from nv_dashboard_widgets import DashboardPage
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'活动.xlsx';pd.DataFrame([fixtures.BASE]).to_excel(p,index=False)
            page=DashboardPage(None,Path(tmp)/'state.db')
            page._on_loaded(fixtures.build_dashboard_snapshot(p,None,page.state_store))
            page.month_mode_combo.setCurrentIndex(page.month_mode_combo.findData('single'))
            pd.DataFrame([{**fixtures.BASE,'活动时间':'2026-08-01'}]).to_excel(p,index=False)
            page._on_loaded(fixtures.build_dashboard_snapshot(p,None,page.state_store))
            self.assertEqual(page.month_start_combo.currentData(),'2026-09')
            self.assertEqual(page.view.activity_count,0);page.close()

    def test_n2_08_actual_chart_labels_have_contrast(self):
        from nv_dashboard_widgets import _HorizontalBarChart,_contrast_ratio
        from PySide6.QtGui import QColor
        chart=_HorizontalBarChart();chart.set_palette({'panel':'#191b20','focus':'#a9bcff'})
        chart.set_data({'甲大学':10},'')
        self.assertGreaterEqual(_contrast_ratio(chart._bar_set.labelColor(),QColor('#a9bcff')),4.5)
        chart.close()


if __name__=='__main__': unittest.main(verbosity=2)
