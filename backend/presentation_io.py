"""Local presentation material extraction and inspection of the exported Office file."""
from __future__ import annotations
import hashlib
import html
import json
from pathlib import Path
import re
import posixpath
import math
import zipfile
import xml.etree.ElementTree as ET


def read_source(args):
    source = Path(args['source']).expanduser().resolve(strict=True)
    if not source.is_file() or source.stat().st_size > 80_000_000:
        raise ValueError('材料须是小于 80MB 的本地文件')
    fragments = []
    extracted_length = 0
    def add(location, text):
        nonlocal extracted_length
        text = str(text).strip()
        extracted_length += len(text)
        if extracted_length > 8_000_000 or len(fragments)>20000:
            raise ValueError('材料超过 800 万字或 2 万个片段，请先按章节/工作表拆分后导入')
        if text:
            for start in range(0, len(text), 8000):
                fragments.append({'location': location + (f' / 字符 {start+1}–{min(start+8000,len(text))}' if len(text)>8000 else ''), 'text': text[start:start+8000]})
    ext = source.suffix.lower()
    if ext == '.docx':
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        doc = Document(source)
        for i, element in enumerate(doc.element.body, 1):
            if element.tag.endswith('}p'):
                add(f'正文块 {i}', Paragraph(element, doc).text)
            elif element.tag.endswith('}tbl'):
                table = Table(element, doc)
                for r, row in enumerate(table.rows, 1):
                    add(f'正文块 {i} / 表格行 {r}', ' | '.join(c.text for c in row.cells))
    elif ext == '.xlsx':
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
        book = load_workbook(source, read_only=True, data_only=False)
        try:
            for sheet in book:
                for i, row in enumerate(sheet.iter_rows(), 1):
                    cells = [f'{get_column_letter(c.column)}{i}={c.value}' for c in row if c.value is not None]
                    if cells: add(f'{sheet.title}!行{i}', ' | '.join(cells))
        finally: book.close()
    elif ext == '.pptx':
        ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
        with zipfile.ZipFile(source) as archive:
            slides = sorted((n for n in archive.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml', n)), key=lambda n: int(re.search(r'slide(\d+)', n)[1]))
            for i, name in enumerate(slides, 1):
                if archive.getinfo(name).file_size > 8_000_000: raise ValueError('PPTX 单页数据过大')
                node = ET.fromstring(archive.read(name))
                for k, para in enumerate(node.findall('.//a:p', ns), 1):
                    add(f'第 {i} 页 / 文本块 {k}', ''.join(t.text or '' for t in para.findall('.//a:t', ns)))
    elif ext == '.pdf':
        from PyPDF2 import PdfReader
        reader = PdfReader(source)
        for i, page in enumerate(reader.pages, 1): add(f'第 {i} 页', page.extract_text() or '')
    elif ext in ('.txt', '.md', '.csv', '.tsv'):
        raw = source.read_bytes()
        try: text = raw.decode('utf-8-sig')
        except UnicodeDecodeError: text = raw.decode('gb18030')
        for i, line in enumerate(text.splitlines(), 1): add(f'行 {i}', line)
    else:
        raise ValueError('材料支持 DOCX、XLSX、PPTX、PDF、TXT、MD、CSV、TSV；旧版 Office 文件请先另存为新格式')
    if not fragments: raise ValueError('没有读取到文字；扫描 PDF 需要先做 OCR')
    offset = int(args.get('offset', 0))
    limit = min(150, max(1, int(args.get('limit', 100))))
    if offset < 0 or offset >= len(fragments): raise ValueError('材料片段 offset 超出范围')
    selected, length = [], 0
    for fragment in fragments[offset:offset+limit]:
        if selected and length + len(fragment['text']) > 150000: break
        selected.append(fragment); length += len(fragment['text'])
    end = offset + len(selected)
    return {'label': source.name, 'original_path': str(source), 'source_hash': hashlib.sha256(source.read_bytes()).hexdigest(),
            'fragments': selected, 'offset': offset, 'total_fragments': len(fragments), 'next_offset': end if end < len(fragments) else None,
            'notes': 'XLSX 保留单元格地址和公式，公式没有在本工具中重新计算。' if ext=='.xlsx' else 'PPTX 读取可见文本与表格；图片、图表及讲者备注需另行解析/视觉审阅，不代表完整读取或无损导入。' if ext=='.pptx' else ''}


def _plain(value):
    return re.sub(r'\s+', '', html.unescape(re.sub(r'<[^>]*>', '', str(value))))


def verify_chart_data(chart, element):
    """Compare ordered OOXML cache values, independently of visual model guesses."""
    ns = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
    errors, limitations = [], []
    series = chart.findall('.//c:ser', ns)
    planned = element['series']
    if any(s['type'] in ('waterfall', 'candlestick') for s in planned):
        return [], ['瀑布/蜡烛图采用派生系列，当前未逐项验证其缓存数值']
    if len(series) != len(planned): return ['导出图表系列数量与数据计划不一致'], []
    for actual, spec in zip(series, planned):
        kind, encode = spec['type'], spec['encode']
        if kind in ('scatter', 'bubble'):
            mappings = [('xVal', 'x', True), ('yVal', 'y', True)] + ([('bubbleSize', 'size', True)] if kind=='bubble' else [])
        else:
            horizontal = kind=='bar' and element.get('yAxis', {}).get('type')=='category'
            mappings = [('cat', 'category' if kind in ('pie','radar') else 'y' if horizontal else 'x', False), ('val', 'value' if kind=='pie' else 'x' if horizontal else 'y', True)]
        for tag, channel, numeric in mappings:
            values = [row[element['data']['cols'].index(encode[channel])] for row in element['data']['rows']]
            container = actual.find('c:'+tag, ns)
            if container is None:
                errors.append(f'缺失图表通道 {tag}'); continue
            points = {int(p.get('idx')): p.findtext('c:v', default='', namespaces=ns) for p in container.findall('.//c:pt', ns)}
            for i, value in enumerate(values):
                observed = points.get(i)
                if value is None and observed in (None, ''): continue
                try: matches = math.isclose(float(observed), float(value), rel_tol=1e-9, abs_tol=1e-9) if numeric else str(observed)==str(value)
                except (ValueError, TypeError): matches = False
                if not matches: errors.append(f'{tag} 第 {i+1} 项不一致：预期 {value}，导出 {observed}')
            if any(i >= len(values) for i in points): errors.append(f'{tag} 含多余数据点')
    return errors, limitations


def verify_presentation(args):
    source = Path(args['source']).resolve(strict=True)
    expected = json.loads(Path(args['expected']).read_text(encoding='utf-8'))
    out = Path(args['output_dir']).resolve(strict=True)
    ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main', 'p': 'http://schemas.openxmlformats.org/presentationml/2006/main', 'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
    report = {'status': 'structure_checked', 'source_sha256': hashlib.sha256(source.read_bytes()).hexdigest(), 'errors': [], 'warnings': [], 'slides': [], 'fonts_embedded': False,
              'limitations': ['自动检查不能替代人工审美判断；交付前仍需查看渲染图片。']}
    with zipfile.ZipFile(source) as archive:
        if archive.testzip(): raise ValueError('PPTX 压缩包校验失败')
        names = archive.namelist()
        slide_files = sorted((n for n in names if re.fullmatch(r'ppt/slides/slide\d+\.xml', n)), key=lambda n: int(re.search(r'slide(\d+)', n)[1]))
        report['fonts_embedded'] = any(n.startswith('ppt/fonts/') for n in names)
        if len(slide_files) != len(expected['pages']): report['errors'].append('导出页数与项目不一致')
        for index, (name, planned) in enumerate(zip(slide_files, expected['pages']), 1):
            node = ET.fromstring(archive.read(name))
            actual = _plain(''.join(t.text or '' for t in node.findall('.//a:t', ns)))
            elements = planned['page']['elements']
            texts = [e.get('content', {}).get('text', '') for e in elements if e['elementType'] in ('text', 'shape')]
            texts += [c.get('text', '') for e in elements if e['elementType']=='table' for row in e['rows'] for c in row]
            missing = [t for t in texts if _plain(t) and _plain(t) not in actual]
            if missing: report['errors'].append(f'第 {index} 页丢失文字/表格内容：' + str(missing[:3])[:250])
            for text in texts:
                paragraphs=[part for part in re.split(r'</p>|<br\s*/?>|\n',str(text),flags=re.I) if _plain(part)]
                if len(paragraphs)<2: continue
                containers=node.findall('.//p:sp',ns)+node.findall('.//a:tc',ns)
                matched=[item for item in containers if _plain(''.join(t.text or '' for t in item.findall('.//a:t',ns)))==_plain(text)]
                if matched and not any(sum(bool(_plain(''.join(t.text or '' for t in para.findall('.//a:t',ns)))) for para in item.findall('.//a:p',ns))>=len(paragraphs) for item in matched):
                    report['errors'].append(f'第 {index} 页段落换行在导出时丢失')
            native = sum(e['elementType']=='chart' and e['series'][0]['type'] not in ('heatmap','treemap','sunburst','sankey') for e in elements)
            charts = len(node.findall('.//c:chart', ns))
            tables = len(node.findall('.//a:tbl', ns))
            pictures = len(node.findall('.//p:pic', ns))
            fallback_count = sum(f['page_id']==planned['id'] for f in expected.get('fallbacks', []))
            if charts != native: report['errors'].append(f'第 {index} 页原生图表数量异常：{charts}/{native}')
            chart_elements = [e for e in elements if e['elementType']=='chart' and e['series'][0]['type'] not in ('heatmap','treemap','sunburst','sankey')]
            if chart_elements:
                rel_name = posixpath.join(posixpath.dirname(name), '_rels', posixpath.basename(name)+'.rels')
                rels = {r.get('Id'): r for r in ET.fromstring(archive.read(rel_name))}
                for ref, element in zip(node.findall('.//c:chart', ns), chart_elements):
                    rel = rels[ref.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')]
                    target = posixpath.normpath(posixpath.join(posixpath.dirname(name), rel.get('Target')))
                    if rel.get('TargetMode')=='External' or not target.startswith('ppt/charts/'):
                        report['errors'].append(f'第 {index} 页图表引用无效'); continue
                    issues, limits = verify_chart_data(ET.fromstring(archive.read(target)), element)
                    report['errors'].extend(f'第 {index} 页 {element["elementId"]}：{issue}' for issue in issues)
                    report['limitations'].extend(limits)
            if tables != sum(e['elementType']=='table' for e in elements): report['errors'].append(f'第 {index} 页原生表格数量异常')
            if pictures < fallback_count: report['errors'].append(f'第 {index} 页复杂图表图片缺失')
            report['slides'].append({'index':index, 'page_id':planned['id'], 'native_charts':charts, 'native_tables':tables, 'pictures':pictures})
    app = deck = None
    existing = False
    try:
        import pythoncom
        import win32com.client
        pythoncom.CoInitialize()
        try:
            try: win32com.client.GetActiveObject('PowerPoint.Application'); existing = True
            except Exception: pass
            app = win32com.client.DispatchEx('PowerPoint.Application')
            deck = app.Presentations.Open(str(source), True, False, False)
            render = out/'render'; render.mkdir(exist_ok=True)
            for item in report['slides']:
                slide = deck.Slides(item['index'])
                png = render/f"slide-{item['index']:02d}.png"
                slide.Export(str(png), 'PNG', 1280, 720)
                item['render'] = str(png)
                item['text_measurements'] = []
                for shape in slide.Shapes:
                    if shape.HasTextFrame and shape.TextFrame.HasText:
                        try:
                            frame = shape.TextFrame2
                            used = frame.TextRange.BoundHeight
                            available = shape.Height - frame.MarginTop - frame.MarginBottom
                            width_available = shape.Width - frame.MarginLeft - frame.MarginRight
                            item['text_measurements'].append({'shape':shape.Name, 'text_excerpt':str(frame.TextRange.Text)[:80],
                                'bounds':[round(float(v),2) for v in (shape.Left,shape.Top,shape.Width,shape.Height)],
                                'used_height':round(float(used),2),'available_height':round(float(available),2),
                                'font':str(frame.TextRange.Font.Name),'font_size':float(frame.TextRange.Font.Size)})
                            if used > available + 5:
                                report['warnings'].append(f"第 {item['index']} 页 {shape.Name} 文字高度接近或超过文本框，请查看渲染")
                            if frame.TextRange.BoundWidth > width_available + 5:
                                report['warnings'].append(f"第 {item['index']} 页 {shape.Name} 文字宽度超过文本框，请查看渲染")
                        except Exception: pass
            report['status'] = 'rendered_with_powerpoint'
            report['render_engine'] = 'Microsoft PowerPoint'
            try:
                pdf = out/'presentation.pdf'
                deck.SaveAs(str(pdf), 32)
                from PyPDF2 import PdfReader
                pdf_pages = len(PdfReader(pdf).pages)
                if pdf_pages != len(expected['pages']):
                    raise ValueError(f'PDF 页数不一致：{pdf_pages}')
                report['pdf'] = str(pdf)
                report['pdf_pages'] = pdf_pages
            except Exception as error:
                report['warnings'].append('PDF 预览导出未完成：' + str(error)[:200])
        finally:
            if deck: deck.Close()
            if app and not existing and app.Presentations.Count == 0: app.Quit()
            pythoncom.CoUninitialize()
    except Exception as error:
        report['warnings'].append('PowerPoint 实际渲染未完成：' + str(error)[:300])
        report['limitations'].append('当前只完成结构检查，不能声明实际显示已通过。')
    if report['status'] == 'rendered_with_powerpoint':
        from PIL import Image, ImageDraw
        columns = min(3, len(report['slides']))
        rows = (len(report['slides']) + columns-1)//columns
        contact = Image.new('RGB', (columns*400, rows*252), '#e4e6e7')
        draw = ImageDraw.Draw(contact)
        for k, item in enumerate(report['slides']):
            with Image.open(item['render']) as opened: thumb = opened.convert('RGB').resize((384,216))
            x, y = (k%columns)*400+8, (k//columns)*252+26
            contact.paste(thumb,(x,y)); draw.text((x,y-20),str(k+1), fill='black')
        contact.save(out/'contact.png'); report['contact_sheet'] = str(out/'contact.png')
    report['report_path'] = str(out/'qa.json')
    (out/'qa.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    return {'presentation_qa': report, 'output_dir': str(out)}
