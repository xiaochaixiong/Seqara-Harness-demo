"""Reproduction and regression of R2-01..27. Physical printers are never used."""
import os
os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
import sys
import tempfile
import threading
import unittest
import zipfile
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pandas as pd
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'native'))
import nv_business_core as biz
import nv_classify_core as invoices
import nv_auto_print_core as printing
from nv_briefing_core import ExcelSplitter, ReportGenerator
from nv_dashboard_core import build_dashboard_snapshot, DashboardStateStore

BASE = {'学校名称':'甲大学','活动名称':'培训','活动类型':'师资培训',
        '活动时间':'2026-09-01','覆盖专业':'设计','服务教师（人数）':10,'服务学生（人数）':20}
APP = {'活动名称':'同名培训','学院名称':'艺术学院','活动类型':'师资培训',
       '开始时间':datetime(2026,9,1,10,30),'服务教师（人数）':10,'服务学生（人数）':20,
       '活动总额':1000,'申请人部门':'驻校办公室/甲大学','审批编号':'202609010001'}


class Round2Repairs(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.home = Path(self.tmp.name)
        self.messages, self.errors = [], []
        for name, fn in [('show_info',lambda *x:self.messages.append(x)),
                         ('show_error',lambda *x:self.errors.append(x)), ('LOG_DATA',str(self.home/'business.log'))]:
            p=patch.object(biz,name,fn);p.start();self.addCleanup(p.stop)
        p=patch.object(invoices,'config_base_path',str(self.home));p.start();self.addCleanup(p.stop)

    def book(self, name, sheets):
        p=self.home/name; w=openpyxl.Workbook();w.remove(w.active)
        for title, rows in sheets.items():
            s=w.create_sheet(title)
            for row in rows:s.append(row)
        w.save(p);w.close();return p

    def frame(self, name, rows):
        p=self.home/name;pd.DataFrame(rows).to_excel(p,index=False);return p

    def read(self,p):
        w=openpyxl.load_workbook(p,data_only=True)
        try:return {s.title:list(s.values) for s in w}
        finally:w.close()

    def invoice(self,p,**kw):
        return invoices.process(p, output_dir=self.home/'invoice-out', company_name='测试公司',
                                group_cover_prefix=kw.pop('project','测试项目'), base=1000,jitter=0,**kw)

    def stat(self,rows,**kw):
        p=self.frame('stat.xlsx',rows)
        result=biz.generate_project_stat_tables(p,self.home/'stat-out',**kw)
        self.assertIsNotNone(result,self.errors)
        return self.read(result)

    def handler(self):
        h=printing._PrintHandler(self.home,None,lambda:None,lambda:False,lambda:'shrink',
            lambda:False,lambda:'word',lambda:False,lambda:'',lambda:False,lambda:False,
            lambda x:self.messages.append(x),threading.Lock(),lambda *x:self.messages.append(x))
        self.addCleanup(h.stop);return h

    def test_r2_01_04_split_merged_rows_and_dimensions(self):
        p=self.book('源.xlsx',{'数据':[['学校名称','活动','备注'],['甲大学','A','甲备注'],
            [None,'B',None],['乙大学','C','乙备注1'],['乙大学','D','乙备注2']]})
        w=openpyxl.load_workbook(p);s=w.active;s.merge_cells('A2:A3');s.merge_cells('C2:C3')
        s.row_dimensions[4].height=80;s.row_dimensions[4].hidden=True;w.save(p);w.close()
        result=ExcelSplitter(lambda _:None).split(p,self.home/'split')
        self.assertEqual(result['success'],2,result)
        self.assertEqual(len(self.read(self.home/'split/甲大学.xlsx')['数据']),3)
        w=openpyxl.load_workbook(self.home/'split/乙大学.xlsx')
        self.assertEqual([w.active['C2'].value,w.active['C3'].value],['乙备注1','乙备注2'])
        self.assertEqual(w.active.row_dimensions[2].height,80)
        self.assertTrue(w.active.row_dimensions[2].hidden);w.close()

    def test_r2_02_duplicate_exclusion_follows_activity(self):
        a={**BASE,'新闻链接':'https://example.invalid/a'}
        b={**BASE,'新闻链接':'https://example.invalid/b','服务教师（人数）':30}
        p=self.frame('同名.xlsx',[a,b]);db=self.home/'state.db'
        s=build_dashboard_snapshot(p,None,db)
        DashboardStateStore(db).set_coverage_excluded(s.activities[0].approval_id,'teacher',True)
        self.frame('同名.xlsx',[b,a]);s=build_dashboard_snapshot(p,None,db)
        self.assertEqual(s.teacher_total,30)
        self.assertEqual(next(a.detail_url for a in s.activities if a.teacher_excluded),a['新闻链接'])

    def test_r2_03_08_09_invoice_totals_multisheet_rerun(self):
        p=self.book('发票.xlsx',{'第一批':[['发票项目','金额（元）'],['住宿费',100],['合计',100]],
            '第二批':[['费用名称','金额'],['住宿费',200],['总计',200]]})
        for _ in range(3):
            p=self.invoice(p)
            rows=self.read(p)['分组统计'];self.assertEqual(rows[1][2],300,rows)

    def test_r2_10_invoice_invalid_long_sheet_name(self):
        p=self.book('发票.xlsx',{'数据':[['发票项目','金额'],['住宿费',100]]})
        output=self.invoice(p,project='2026/9 项目：'+('长名称'*20))
        w=openpyxl.load_workbook(output)
        self.assertTrue(all(len(s.title)<=31 for s in w));w.close()

    def test_r2_11_statistics_comma_and_negative_counts(self):
        tables=self.stat([{**BASE,'服务教师（人数）':'1,200'}])
        self.assertEqual(tables['表4_专业覆盖(全额)'][1][-1],1200)
        p=self.frame('negative.xlsx',[{**BASE,'服务学生（人数）':-20}])
        self.assertIsNone(biz.generate_project_stat_tables(p,self.home/'negative'))
        self.assertIn('非负整数',str(self.errors))

    def test_r2_12_statistics_allocation_conserves_total(self):
        rows=[{**BASE,'覆盖专业':'甲、乙、丙','服务教师（人数）':1,'服务学生（人数）':1} for _ in range(100)]
        table=self.stat(rows)['表5_专业覆盖(均分)']
        self.assertAlmostEqual(table[-1][-1],100,places=8)
        self.assertAlmostEqual(sum(r[-1] for r in table[1:-1]),100,places=8)

    def test_r2_13_20_statistics_duplicate_and_custom_labels(self):
        tables=self.stat([{**BASE,'覆盖专业':'设计、设计','活动类型':'师资培训、师资培训'}],activity_types=['师资培训','未分类'])
        self.assertEqual(tables['表4_专业覆盖(全额)'][1][2:],(1,20,10))
        self.assertEqual(tables['表3_类型分布统计'][1][-1],1)

    def test_r2_14_15_16_clean_dates_ids_and_multiple_links(self):
        p=self.frame('申请.xlsx',[APP,{**APP,'审批编号':'202609010002'}])
        r=self.frame('报销.xlsx',[{'活动名称':'合并报销','关联申请单':'202609010001,202609010002',
                                '新闻链接':'https://example.invalid/both'}])
        result=biz.match_and_clean(p,r,self.home/'clean',None)
        self.assertIsNotNone(result,self.errors)
        df=pd.read_excel(result,dtype={'审批编号':str})
        self.assertEqual(len(df),2)
        self.assertEqual(set(df['审批编号']),{'202609010001','202609010002'})
        self.assertTrue(df['活动时间'].notna().all())
        self.assertTrue(df['新闻链接'].str.endswith('/both').all())

    def test_r2_17_18_briefing_blank_name_and_header_offset(self):
        p=self.frame('简报.xlsx',[{'学校名称':'甲大学','活动名称':None,'活动内容':'本地正文'}])
        result=ReportGenerator(lambda _:None,row_delay_sec=0,split_by_activity=True).generate(p,self.home/'brief')
        self.assertTrue(result['artifacts'],result)
        text='\n'.join(p.text for p in Document(result['artifacts'][0]).paragraphs)
        self.assertIn('未命名活动',text);self.assertNotIn('nan',text)
        p=self.book('表头.xlsx',{'数据':[['总标题'],['学校名称','活动名称','活动内容'],[None,'活动','正文']]})
        result=ReportGenerator(lambda _:None,row_delay_sec=0).generate(p,self.home/'header')
        self.assertIn('3',str(result));self.assertEqual(result['success'],0)

    def test_r2_19_source_overwrite_protected(self):
        p=self.frame('申请.xlsx',[APP]);r=self.frame('报销.xlsx',[{'活动名称':'合并报销','关联申请单':'202609010001','新闻链接':''}])
        before=p.read_bytes()
        self.assertIsNone(biz.match_and_clean(p,r,self.home,None,output_basename=p.stem))
        self.assertEqual(p.read_bytes(),before)
        p=self.frame('统计.xlsx',[BASE]);before=p.read_bytes()
        self.assertIsNone(biz.generate_project_stat_tables(p,self.home,output_filename=p.name))
        self.assertEqual(p.read_bytes(),before)

    def test_r2_21_disable_invoice_amount_fallback(self):
        p=self.book('发票.xlsx',{'数据':[['发票项目','金额'],['ZXQ测试项目',100]]})
        out=self.invoice(p,enable_amount_fallback=False)
        self.assertIn('未分类',str(self.read(out)['分类&分组']))

    def test_r2_22_zip_retry_only_failed_files(self):
        p=self.home/'documents.zip'
        with zipfile.ZipFile(p,'w') as z:z.writestr('A.pdf',b'A');z.writestr('B.pdf',b'B')
        h=self.handler();calls=[]
        def dispatch(p,*_):
            calls.append(p.name);return p.name!='B.pdf' or calls.count('B.pdf')>1
        with patch.object(printing,'wait_file_size_stable',return_value=True),patch.object(printing,'dispatch_print',side_effect=dispatch),patch.object(h,'_schedule'):
            h._try_print(str(p.resolve()));h._try_print(str(p.resolve()))
        self.assertEqual(calls,['A.pdf','B.pdf','B.pdf'])

    def test_r2_23_download_renamed_to_pdf(self):
        from watchdog.events import FileMovedEvent
        h=self.handler()
        with patch.object(h,'_schedule') as schedule:
            h.dispatch(FileMovedEvent(str(self.home/'download.crdownload'),str(self.home/'invoice.pdf')))
            schedule.assert_called_once_with(self.home/'invoice.pdf')

    def test_r2_24_empty_file_wait_is_bounded(self):
        p=self.home/'empty.pdf';p.write_bytes(b'');h=self.handler()
        with patch.object(h,'_schedule') as schedule,patch.object(printing.time,'monotonic',side_effect=[0,0,31,31]):
            h._try_print(str(p));h._try_print(str(p))
            self.assertEqual(schedule.call_count,1)
        self.assertIn(('failed',str(p)),self.messages)

    def test_r2_27_statistics_message_matches_artifact(self):
        tables=self.stat([{'学校名称':'甲大学','活动类型':'师资培训'}])
        self.assertEqual(len(tables),3)
        self.assertIn('3',str(self.messages))
        self.assertNotIn('表4',str(self.messages))


class Round2WidgetRepairs(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from PySide6.QtWidgets import QApplication
        cls.app=QApplication.instance() or QApplication([])

    def test_r2_07_month_survives_refresh(self):
        from nv_dashboard_widgets import DashboardPage
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'data.xlsx'
            pd.DataFrame([BASE,{**BASE,'活动时间':'2026-08-01'}]).to_excel(p,index=False)
            page=DashboardPage(None,Path(tmp)/'state.db')
            s=build_dashboard_snapshot(p,None,page.state_store)
            page._on_loaded(s)
            page.month_mode_combo.setCurrentIndex(page.month_mode_combo.findData('single'))
            page.month_start_combo.setCurrentIndex(page.month_start_combo.findData('2026-09'))
            page._on_loaded(build_dashboard_snapshot(p,None,page.state_store))
            self.assertEqual(page.month_mode_combo.currentData(),'single')
            self.assertEqual(page.month_start_combo.currentData(),'2026-09')
            self.assertEqual(page.view.activity_count,1);page.close()

    def test_r2_25_26_dark_labels_and_full_tooltip(self):
        import json
        from PySide6.QtGui import QColor
        from nv_dashboard_widgets import _HeatmapTable,_HorizontalBarChart,_composite_color,_contrast_ratio,_table_item
        tokens=json.loads((ROOT/'native/ui_tokens.json').read_text(encoding='utf8'))['dark']
        p={'panel':tokens['canvas'],'focus':tokens['focus'],'text':tokens['text'],'muted':tokens['text_secondary'] if 'text_secondary' in tokens else tokens['text']}
        table=_HeatmapTable();table.set_palette(p);table.set_matrix(['大学'],['月份'],{'大学':{'月份':1}})
        item=table.item(0,0)
        self.assertGreaterEqual(_contrast_ratio(item.foreground().color(),_composite_color(item.background().color(),QColor(p['panel']))),4.5)
        self.assertGreaterEqual(_contrast_ratio(QColor('#202328'),QColor(p['focus'])),4.5)
        long='长活动名称'*40;self.assertEqual(_table_item(long).toolTip(),long);table.close()


if __name__=='__main__': unittest.main(verbosity=2)
