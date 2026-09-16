"""End-to-end worker regressions; NV_FIX_WORKER_PATH selects the packaged EXE."""
import os,sys,unittest
from pathlib import Path
import pandas as pd
import openpyxl
import test_round2_repairs as fixtures
import test_user_audit_repairs as old


class ReleaseScenarios(unittest.TestCase):
    setUp=fixtures.Round2Repairs.setUp
    book=fixtures.Round2Repairs.book
    frame=fixtures.Round2Repairs.frame
    read=fixtures.Round2Repairs.read
    invoke=old.UserAuditRepairs.invoke

    def excel(self, event):
        return next(p for p in event['artifacts'] if p.endswith('.xlsx'))

    def test_release_invoice_multiple_sheets_footer_and_repeat(self):
        p=self.book('发票.xlsx',{'一期':[['序号','发票项目','金额'],[1,'住宿费',100],['合计',None,100]],
            '二期':[['发票项目','金额'],['住宿费',200]]})
        args={'source':str(p),'company':'测试公司','project':'2026/9 '+('长项目'*12),'base':1000,'jitter':0}
        output=self.excel(self.invoke('invoices',args));self.assertEqual(self.read(output)['分组统计'][1][2],300)
        output=self.excel(self.invoke('invoices',{**args,'source':output}))
        self.assertEqual(self.read(output)['分组统计'][1][2],300)

    def test_release_clean_distinct_approvals_and_date(self):
        p=self.frame('申请.xlsx',[fixtures.APP,{**fixtures.APP,'审批编号':'202609010002'}])
        r=self.frame('报销.xlsx',[{'活动名称':'合并报销','关联申请单':'202609010001,202609010002','新闻链接':'https://example.invalid/news'}])
        out=self.excel(self.invoke('clean',{'source':str(p),'reimbursement':str(r)}))
        df=pd.read_excel(out);self.assertEqual(len(df),2)
        self.assertTrue(df['活动时间'].notna().all());self.assertTrue(df['新闻链接'].notna().all())

    def test_release_statistics_headcounts_and_unknown_labels(self):
        p=self.frame('统计.xlsx',[{**fixtures.BASE,'服务教师（人数）':'1,200','覆盖专业':'设计、设计','活动类型':'未知甲、未知乙'}])
        tables=self.read(self.excel(self.invoke('statistics',{'source':str(p)})))
        self.assertEqual(tables['表4_专业覆盖(全额)'][1][-1],1200)
        self.assertEqual(tables['表3_类型分布统计'][1][-1],1)
        p=self.frame('异常.xlsx',[{**fixtures.BASE,'服务教师（人数）':'1,2'}])
        self.invoke('statistics',{'source':str(p)},False)

    def test_release_split_merged_school_and_row_height(self):
        p=self.book('拆分.xlsx',{'数据':[['学校名称','活动'],['甲大学','A'],[None,'B'],['乙大学','C']]})
        w=openpyxl.load_workbook(p);w.active.merge_cells('A2:A3');w.active.row_dimensions[4].height=80;w.save(p);w.close()
        event=self.invoke('split',{'source':str(p)})
        names={Path(p).name:p for p in event['artifacts']}
        self.assertEqual(len(self.read(names['甲大学.xlsx'])['数据']),3)
        w=openpyxl.load_workbook(names['乙大学.xlsx']);self.assertEqual(w.active.row_dimensions[2].height,80);w.close()

    def test_release_briefing_blank_and_long_titles(self):
        p=self.frame('简报.xlsx',[{'学校名称':'甲大学','活动名称':None,'活动内容':'本地正文'},
            {'学校名称':'甲大学','活动名称':'长活动名称'*70,'活动内容':'完整内容'}])
        event=self.invoke('briefing',{'source':str(p),'split_by_activity':True})
        self.assertEqual(event['briefing']['status'],'completed')
        self.assertEqual(event['briefing']['success'],2)
        self.assertEqual(len(event['artifacts']),2)


if __name__=='__main__':unittest.main(verbosity=2)
