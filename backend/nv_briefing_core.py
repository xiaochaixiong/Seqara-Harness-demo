# -*- coding: utf-8 -*-
"""
活动简报：Excel 按校拆分 + 抓取链接生成 Word。
从「工具集V4.3 CTK替换新UI底层.py」抽出，供 PySide6 版调用（避免加载 Tk/CTK）。

抓取强化要点：现代浏览器 UA、Accept-Language、Referer；HTML 解码兜底；
官网优先 <main> 与 JSON-LD articleBody/description；元数据含 twitter:description；
微信页仍受验证/风控限制，不强行用 JSON-LD 覆盖。
"""
from __future__ import annotations
from nv_data_safety import safe_file_stem, staged_outputs, parse_headcount

import json
import os
import re
import shutil
import time
import html
import uuid
from datetime import datetime
from io import BytesIO
from copy import copy
from pathlib import Path
from typing import Any, Callable, List, Optional
from urllib.parse import urljoin, urlparse

import openpyxl
from openpyxl.formula.tokenizer import Tokenizer
import pandas as pd
import requests
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

try:
    from urllib3.exceptions import InsecureRequestWarning

    import urllib3

    urllib3.disable_warnings(InsecureRequestWarning)
except Exception:
    pass

try:
    from bs4 import BeautifulSoup
except ImportError as exc:
    raise RuntimeError("缺少 beautifulsoup4；请在构建或部署阶段安装完整依赖。") from exc

try:
    from PIL import Image
    from PIL import ImageStat
except ImportError as exc:
    raise RuntimeError("缺少 Pillow；请在构建或部署阶段安装完整依赖。") from exc


class ExcelSplitter:
    def __init__(self, log_callback, dialog_callback=None):
        self.log = log_callback
        self.dialog_callback = dialog_callback
        self.stop_flag = False

    def _dialog(self, kind, title, message):
        if self.dialog_callback:
            self.dialog_callback(kind, title, message)

    def split(self, master_path, output_dir):
        master = Path(master_path)
        output = Path(output_dir)
        self.log(f"正在分析总表: {master} ...")
        result = {"total": 0, "success": 0, "failed": []}
        if master.suffix.lower() != ".xlsx":
            msg = "表格拆分仅支持 .xlsx；.xls 或 .csv 请先另存为 .xlsx。"
            self.log(f"❌ {msg}")
            self._dialog("error", "格式不支持", msg)
            return result
        try:
            try:
                source_bytes = master.read_bytes()
                probe = openpyxl.load_workbook(BytesIO(source_bytes), read_only=True, data_only=True)
            except Exception as e:
                msg = f"读取 Excel 失败，请确认文件未损坏且未加密：{e}"
                self.log(f"❌ {msg}")
                self._dialog("error", "读取失败", msg)
                return result

            source_sheet = ""
            school_col_index = 0
            header_row_index = 0
            schools: list[str] = []
            try:
                for ws in probe.worksheets:
                    school_col_index = 0
                    for r in range(1, min(ws.max_row, 12) + 1):
                        for c in range(1, ws.max_column + 1):
                            if str(ws.cell(row=r, column=c).value or "").strip() == "学校名称":
                                source_sheet = ws.title
                                school_col_index = c
                                header_row_index = r
                                break
                        if school_col_index:
                            break
                    if school_col_index:
                        for r in range(header_row_index + 1, ws.max_row + 1):
                            value = str(ws.cell(row=r, column=school_col_index).value or "").strip()
                            if value and value not in schools:
                                schools.append(value)
                        # Other worksheets may contain additional schools.
            finally:
                probe.close()

            if not source_sheet:
                msg = "在工作簿前 12 行中找不到“学校名称”表头。"
                self.log(f"❌ {msg}")
                self._dialog("error", "无法拆分", msg)
                return result
            if not schools:
                msg = "“学校名称”列中没有可拆分的数据。"
                self.log(f"❌ {msg}")
                self._dialog("error", "无法拆分", msg)
                return result

            output.mkdir(parents=True, exist_ok=True)
            total_schools = len(schools)
            result["total"] = total_schools
            self.log(f"🔍 检测到 {total_schools} 所学校，准备开始拆分...")
            used_stems: set[str] = set()
            for idx, school in enumerate(schools):
                if self.stop_flag:
                    self.log("任务已停止")
                    break
                school_name = str(school).strip()
                safe_name = safe_file_stem(school_name, f"学校_{idx + 1:03d}")
                stem = safe_name
                duplicate_index = 2
                while stem.casefold() in used_stems or (output / f"{stem}.xlsx").exists():
                    stem = f"{safe_name}_{duplicate_index}"
                    duplicate_index += 1
                safe_name = stem
                used_stems.add(safe_name.casefold())
                self.log(f"[{idx + 1}/{total_schools}] 正在生成：{safe_name}.xlsx ...")
                target_file = output / f"{safe_name}.xlsx"
                temp_file = output / f".{safe_name}.{uuid.uuid4().hex}.tmp.xlsx"
                wb = None
                try:
                    wb = openpyxl.load_workbook(BytesIO(source_bytes))
                    values = openpyxl.load_workbook(BytesIO(source_bytes), data_only=True)
                    filtered_sheets = 0
                    for ws in list(wb.worksheets):
                        local_header = 0
                        local_school_col = 0
                        for r in range(1, min(ws.max_row, 12) + 1):
                            for c in range(1, ws.max_column + 1):
                                if str(ws.cell(row=r, column=c).value or "").strip() == "学校名称":
                                    local_header = r
                                    local_school_col = c
                                    break
                            if local_school_col:
                                break
                        if not local_school_col:
                            self.log(f"未识别学校列，拆分副本不包含工作表：{ws.title}；原总表保留。")
                            wb.remove(ws)
                            continue
                        filtered_sheets += 1
                        # Flatten data merges before deleting rows. Header merges remain.
                        for area in list(ws.merged_cells.ranges):
                            if area.max_row <= local_header:
                                continue
                            anchor = ws.cell(area.min_row, area.min_col)
                            value, style = anchor.value, copy(anchor._style)
                            cached = values[ws.title].cell(area.min_row, area.min_col).value
                            ws.unmerge_cells(str(area))
                            for rr in range(area.min_row, area.max_row + 1):
                                for cc in range(area.min_col, area.max_col + 1):
                                    ws.cell(rr, cc).value = value
                                    ws.cell(rr, cc)._style = copy(style)
                                    # Cached merged values are used only for formula handling.
                                    if values[ws.title].cell(rr, cc).__class__.__name__ != 'MergedCell':
                                        values[ws.title].cell(rr, cc).value = cached
                        kept = [r for r in range(1, ws.max_row + 1) if r <= local_header or
                                str(ws.cell(r, local_school_col).value or '').strip() == school_name]
                        moves = {r: i + 1 for i, r in enumerate(kept)}
                        dimensions = {moves[r]: copy(ws.row_dimensions[r]) for r in kept if r in ws.row_dimensions}
                        for row in kept:
                            for cell in ws[row]:
                                if cell.data_type != 'f':
                                    continue
                                tokens = Tokenizer(cell.value).items
                                simple = all(t.subtype != 'RANGE' or
                                    all(re.fullmatch(r'\$?[A-Za-z]{1,3}\$?' + str(row), ref)
                                        for ref in t.value.split(':')) for t in tokens)
                                if simple:
                                    for token in tokens:
                                        if token.subtype == 'RANGE':
                                            token.value = re.sub(r'(\$?[A-Za-z]{1,3}\$?)\d+',
                                                lambda m: m[1] + str(moves[row]), token.value)
                                    cell.value = '=' + ''.join(t.value for t in tokens)
                                else:
                                    cached = values[ws.title][cell.coordinate].value
                                    if cached is None:
                                        raise ValueError(f'{ws.title}!{cell.coordinate} 的跨行／跨表公式没有缓存结果，请先用 Excel/WPS 重算并保存后再拆分。')
                                    cell.value = cached
                                    self.log(f'{ws.title}!{cell.coordinate} 跨行／跨表公式已保留原计算值。')
                        for row in range(ws.max_row, local_header, -1):
                            cell_val = ws.cell(row=row, column=local_school_col).value
                            if str(cell_val or "").strip() != school_name:
                                ws.delete_rows(row)
                        ws.row_dimensions.clear()
                        for new_row, dimension in dimensions.items():
                            dimension.index = new_row
                            ws.row_dimensions[new_row] = dimension
                    if not filtered_sheets:
                        raise RuntimeError("生成副本时未找到可过滤的“学校名称”数据页")
                    wb.save(temp_file)
                    wb.close()
                    wb = None
                    # Windows rename refuses a target created concurrently; never overwrite.
                    os.rename(temp_file, target_file)
                    result["success"] += 1
                except Exception as e:
                    self.log(f"  ❌ 处理Excel数据失败: {e}")
                    result["failed"].append(f"{school_name}: {e}")
                finally:
                    if wb is not None:
                        try:
                            wb.close()
                        except Exception:
                            pass
                    if 'values' in locals():
                        values.close()
                    try:
                        temp_file.unlink(missing_ok=True)
                    except OSError:
                        pass

            success = int(result["success"])
            failed = total_schools - success
            if failed:
                self.log(f"\n⚠️ 拆分结束：成功 {success} 个，失败或取消 {failed} 个。")
                self._dialog(
                    "error" if success == 0 else "info",
                    "拆分结果",
                    f"成功生成 {success}/{total_schools} 个学校表格；其余文件未发布。\n保存在：{output}",
                )
            else:
                self.log(f"\n✅ 拆分完成！实际生成 {success} 个文件。")
                self._dialog("info", "完成", f"成功拆分出 {success} 个表格！\n保存在：{output}")
        except Exception as e:
            self.log(f"❌ 发生未知错误: {e}")
            result["failed"].append(str(e))
            self._dialog("error", "拆分失败", str(e))
        return result


class ReportGenerator:
    def __init__(
        self,
        log_callback: Callable[[str], None],
        dialog_callback: Optional[Callable[[str, str, str], None]] = None,
        *,
        fetch_timeout: int = 18,
        row_delay_sec: float = 0.45,
        summary_max_chars: int = 420,
        max_images_per_article: int = 3,
        nearby_images_only: bool = True,
        nearby_scan_nodes: int = 280,
        keep_activity_meta: bool = False,
        full_content_mode: bool = False,
        full_content_max_chars: int = 60000,
        full_mode_max_images: int = 40,
        split_by_activity: bool = False,
    ):
        self.log = log_callback
        self.dialog_callback = dialog_callback
        self.stop_flag = False
        self.fetch_timeout = max(5, int(fetch_timeout))
        self.row_delay_sec = max(0.0, float(row_delay_sec))
        self.summary_max_chars = max(120, int(summary_max_chars))
        self.max_images_per_article = max(0, int(max_images_per_article))
        # 仅抓取“正文锚点附近”的图片，减少页眉页脚装饰图混入
        self.nearby_images_only = bool(nearby_images_only)
        self.nearby_scan_nodes = max(80, int(nearby_scan_nodes))
        # True：写入表格中的活动时间、服务教师/学生人数与文末「新闻链接」行；False：仅标题+正文+配图
        self.keep_activity_meta = bool(keep_activity_meta)
        # True：每条链接尽量写入全文 + 正文区内更多配图，输出「全量报告」（体积大、耗时长）
        self.full_content_mode = bool(full_content_mode)
        self.full_content_max_chars = max(8000, min(200000, int(full_content_max_chars)))
        self.full_mode_max_images = max(0, min(80, int(full_mode_max_images)))
        # True：每个活动单独生成一个 Word，文件名与活动名称一致
        self.split_by_activity = bool(split_by_activity)
        self.session = requests.Session()
        retries = Retry(total=3, backoff_factor=1, status_forcelist=[500, 502, 503, 504])
        self.session.mount("http://", HTTPAdapter(max_retries=retries))
        self.session.mount("https://", HTTPAdapter(max_retries=retries))
        self.session.headers.update(
            {
                "User-Agent": (
                    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                    "(KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36"
                ),
                "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
                "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.5",
            }
        )
        self.session.trust_env = False

    def _dialog(self, kind, title, message):
        if self.dialog_callback:
            self.dialog_callback(kind, title, message)

    def set_font(self, run, font_name="仿宋", size=12, bold=False):
        run.font.name = font_name
        try:
            run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
        except Exception:
            pass
        run.font.size = Pt(size)
        run.bold = bold

    def smart_clean_time(self, val):
        if pd.isna(val) or str(val).lower() == "nan":
            return ""
        if isinstance(val, (pd.Timestamp, datetime)):
            return val.strftime("%Y-%m-%d")
        val_str = str(val).strip()
        val_str = re.sub(r"\s+00:00:00(\.0+)?$", "", val_str)
        return val_str

    def extract_url(self, text):
        if not text or str(text).lower() == "nan":
            return None, "内容为空"
        text = str(text).strip()
        pattern = r"https?://[^\s<>'\"，。；;）)\]]+"
        matches = re.findall(pattern, text)
        if matches:
            cleaned = [m.rstrip(".,;:!?)]}，。；：！？") for m in matches if m]
            # 一个单元格里常有多个链接，优先微信公众号正文链接，其次第一个有效链接。
            for u in cleaned:
                if "mp.weixin.qq.com" in u:
                    return u, "成功提取(微信优先)"
            return cleaned[0], "成功提取"
        if re.search(r"[a-zA-Z0-9\.\-]+\.(com|cn|net|org)", text) and not re.search(r"[\u4e00-\u9fa5]", text):
            return "https://" + text, "自动补全https"
        return None, f"未发现有效网址 (原文: {text[:15]}...)"

    def normalize_article_text(self, text):
        t = html.unescape((text or "").replace("\xa0", " "))
        t = re.sub(r"\s+", " ", t).strip()
        # 清理页面上常见的视觉噪音和宣传残留。
        t = re.sub(r"(✦|•|◆|●|■|/){3,}", " ", t)
        t = re.sub(r"左右滑动[，,]?\s*查看更多", " ", t)
        t = re.sub(r"点击.*?阅读原文", " ", t)
        return re.sub(r"\s{2,}", " ", t).strip()

    def _dedup_sentences(self, text):
        parts = [s.strip() for s in re.split(r"(?<=[。！？!?])", text or "") if s.strip()]
        out, seen = [], set()
        for s in parts:
            key = re.sub(r"\W+", "", s)
            if len(key) < 8:
                continue
            if key in seen:
                continue
            seen.add(key)
            out.append(s)
        return "".join(out).strip()

    def _extract_meta_text(self, soup):
        cands = []
        for key, val in [
            ("property", "og:description"),
            ("name", "description"),
            ("name", "twitter:description"),
            ("property", "twitter:description"),
            ("property", "og:title"),
        ]:
            node = soup.find("meta", attrs={key: val})
            if node and node.get("content"):
                cands.append(self.normalize_article_text(node.get("content")))
        cands = [x for x in cands if x and len(x) >= 8]
        return cands[0] if cands else ""

    def _extract_json_ld_article_text(self, soup) -> str:
        """部分高校/政务站用 JSON-LD 标注正文，结构化抽取可显著优于纯 div 猜测。"""
        best = ""
        best_len = 0
        for script in soup.find_all("script"):
            st = (script.get("type") or "").lower()
            if "ld+json" not in st:
                continue
            raw = (script.string or script.get_text() or "").strip()
            if not raw:
                continue
            try:
                data = json.loads(raw)
            except Exception:
                continue

            def walk(o: Any) -> None:
                nonlocal best, best_len
                if isinstance(o, dict):
                    for k in ("articleBody", "description"):
                        v = o.get(k)
                        if isinstance(v, str) and len(v.strip()) > 36:
                            piece = self.normalize_article_text(v)
                            if len(piece) > best_len and self._score_content_quality(piece) >= 18:
                                best, best_len = piece, len(piece)
                    for v in o.values():
                        walk(v)
                elif isinstance(o, list):
                    for v in o:
                        walk(v)

            walk(data)
        return best

    def _decode_response_html(self, response: requests.Response) -> str:
        """尽量正确解码 HTML，减少乱码导致的正文抽取失败。"""
        enc = (getattr(response, "encoding", None) or "").strip()
        if enc and enc.lower() not in ("iso-8859-1",):
            try:
                return response.content.decode(enc, errors="replace")
            except (LookupError, ValueError, TypeError):
                pass
        ct = (response.headers.get("Content-Type") or "").lower()
        m = re.search(r"charset=([\w\-]+)", ct)
        if m:
            try:
                return response.content.decode(m.group(1).strip("'\""), errors="replace")
            except (LookupError, ValueError):
                pass
        apparent = getattr(response, "apparent_encoding", None) or "utf-8"
        try:
            return response.content.decode(apparent, errors="replace")
        except (LookupError, ValueError):
            return response.content.decode("utf-8", errors="replace")

    def _score_content_quality(self, text):
        t = self.normalize_article_text(text)
        if not t:
            return 0
        total = len(t)
        zh = len(re.findall(r"[\u4e00-\u9fff]", t))
        alpha = len(re.findall(r"[A-Za-z]", t))
        symbol = len(re.findall(r"[^\w\u4e00-\u9fff\s]", t))
        score = 0
        if total >= 80:
            score += 30
        if zh / max(total, 1) >= 0.35:
            score += 35
        if alpha / max(total, 1) <= 0.35:
            score += 15
        if symbol / max(total, 1) <= 0.15:
            score += 10
        if re.search(r"环境异常|去验证|轻点两下取消赞", t):
            score -= 50
        if "Promote teaching through competition" in t:
            score -= 15
        quoted = re.findall(r"[「“](.{8,80}?)[”」]", t)
        if quoted and len(set(quoted)) < len(quoted):
            score -= 25
        return score

    def _looks_broken_wechat_text(self, text):
        t = self.normalize_article_text(text)
        if not t:
            return True
        if re.search(r"环境异常|去验证|轻点两下取消赞", t):
            return True
        if "Promote teaching through competition" in t:
            return True
        quoted = re.findall(r"[「“](.{8,80}?)[”」]", t)
        if quoted and len(set(quoted)) < len(quoted):
            return True
        return False

    def _collect_text_blocks(self, root, tags=None, max_blocks=8, max_chars=1200):
        tags = tags or ["p", "div", "span", "section", "font"]
        blocks = []
        seen = set()
        anchor = None
        for node in root.find_all(tags):
            if node.name in ["script", "style"] or node.find_parent("a"):
                continue
            text = self.normalize_article_text(node.get_text(" ", strip=True))
            if not text or self.is_junk_paragraph(text):
                continue
            if text in seen:
                continue
            seen.add(text)
            blocks.append(text)
            if anchor is None:
                anchor = node
            if len("".join(blocks)) >= max_chars or len(blocks) >= max_blocks:
                break
        return blocks, anchor

    def _is_decorative_img_tag(self, img_tag, full_src):
        """根据标签属性和链接特征过滤页面装饰图。

        注意：勿对整段 URL 粗暴匹配 ``thumb``，否则易误杀微信/图床带 ``thumb`` 参数的正文大图。
        """
        attr_s = " ".join(
            [
                str(img_tag.get("class", "")),
                str(img_tag.get("id", "")),
                str(img_tag.get("alt", "")),
                str(img_tag.get("title", "")),
            ]
        ).lower()
        bad_in_attrs = [
            "logo",
            "icon",
            "avatar",
            "qrcode",
            "qr",
            "banner",
            "bg",
            "background",
            "divider",
            "split",
            "decor",
            "ad",
            "loading",
            "watermark",
            "ornament",
            "emoji",
            "face",
            "smiley",
        ]
        if any(k in attr_s for k in bad_in_attrs):
            return True
        if re.search(r"\bthumb\b|thumbnail|_thumb\.|/thumb/", attr_s):
            return True

        url = (full_src or "").lower()
        url_bad_fragments = [
            "logo",
            "favicon",
            "icon.",
            "/icon/",
            "qrcode",
            "qr_code",
            "avatar",
            "1x1",
            "spacer",
            "blank.gif",
            "pixel.gif",
        ]
        if any(k in url for k in url_bad_fragments):
            return True
        return False

    @staticmethod
    def _img_url_from_tag(img_tag) -> Optional[str]:
        """懒加载站点多用 data-src / data-original；兼容微信与常见 CMS。"""
        for key in (
            "data-src",
            "data-original",
            "data-lazy-src",
            "data-lazyload",
            "data-url",
            "data-img",
            "src",
        ):
            v = img_tag.get(key)
            if not v:
                continue
            s = str(v).strip()
            if not s or s.lower().startswith("data:"):
                continue
            return s
        return None

    def _img_under_content_root(self, img_tag, content_div) -> bool:
        if content_div is None or img_tag is None:
            return False
        p = img_tag
        while p is not None:
            if p is content_div:
                return True
            p = getattr(p, "parent", None)
        return False

    def _tag_marker_blob(self, tag) -> str:
        if tag is None or not hasattr(tag, "get"):
            return ""
        parts = [str(tag.get("id", "")), str(tag.get("role", ""))]
        cls = tag.get("class", [])
        if isinstance(cls, list):
            parts.append(" ".join(cls))
        else:
            parts.append(str(cls))
        return " ".join(parts).lower()

    def _tag_looks_like_page_chrome(self, tag) -> bool:
        """侧栏、页眉页脚、导航、推荐区等，不作为正文根容器。"""
        blob = self._tag_marker_blob(tag)
        if not blob.strip():
            return False
        return bool(
            re.search(
                r"\b(header|footer|nav|navbar|nav_|topnav|menu|toolbar|breadcrumb|bread-?crumb|sidebar|side-?bar|"
                r"aside|widget|comment|reply|related|recommend|share|popup|modal|copyright|partner|friend-?link|"
                r"fixed-?(top|bottom)|toolbox|login|register|search-?box|hot-?news|rank|tag-?cloud|"
                r"wechat|qrcode|qr-?code)\b",
                blob,
                re.I,
            )
        )

    def _best_article_subtree(self, soup) -> Any:
        """选取最像单篇新闻正文的 DOM 根，避免整站 layout。"""
        precise_selectors = [
            "#vsb_content",
            "#articleContent",
            ".v_news_content",
            ".TRS_Editor",
            ".article-content",
            ".articleContent",
            ".news_content",
            ".entry-content",
            ".post-content",
            ".detail_content",
            ".news-detail",
            ".n_detail",
            "#BodyLabel",
            ".show_content",
            ".view_content",
        ]
        for sel in precise_selectors:
            for node in soup.select(sel):
                if self._tag_looks_like_page_chrome(node):
                    continue
                txt = (node.get_text() or "").strip()
                if len(txt) < 80:
                    continue
                head = txt[:1200]
                if self._looks_like_nav_text(head):
                    continue
                if self._score_content_quality(head) < 18:
                    continue
                return node

        for art in soup.find_all("article"):
            if self._tag_looks_like_page_chrome(art):
                continue
            txt = (art.get_text() or "").strip()
            if len(txt) < 80:
                continue
            head = txt[:1200]
            if self._looks_like_nav_text(head):
                continue
            if self._score_content_quality(head) < 18:
                continue
            return art

        main_el = soup.find("main")
        if main_el and not self._tag_looks_like_page_chrome(main_el):
            inner = main_el.find("article")
            if inner and not self._tag_looks_like_page_chrome(inner):
                itxt = (inner.get_text() or "").strip()
                if len(itxt) >= 60 and not self._looks_like_nav_text(itxt[:1000]):
                    return inner
            for sel in (".v_news_content", ".article-content", "#vsb_content", ".TRS_Editor", ".news_content"):
                hit = main_el.select_one(sel)
                if hit and not self._tag_looks_like_page_chrome(hit):
                    ht = (hit.get_text() or "").strip()
                    if len(ht) >= 60 and not self._looks_like_nav_text(ht[:1000]):
                        return hit
            mtxt = (main_el.get_text() or "").strip()
            if len(mtxt) >= 120 and not self._looks_like_nav_text(mtxt[:1500]):
                if self._score_content_quality(mtxt[:2000]) >= 22:
                    return main_el

        return None

    def _resolve_content_div(self, soup, host: str, *, full_mode: bool):
        """确定配图与正文块抽取所用的根节点；全量模式禁止退化为整页 body。"""
        if "mp.weixin.qq.com" in host:
            return soup.find(id="js_content") or soup.find("div", class_=re.compile(r"rich_media_content", re.I))

        picked = self._best_article_subtree(soup)
        if picked is not None:
            return picked

        keyword_re = re.compile(r"content|article|post|rich_media|detail|news|entry|vsb|trs|bodylabel", re.I)
        for tag in soup.find_all(["article", "div", "section"]):
            if self._tag_looks_like_page_chrome(tag):
                continue
            sid = str(tag.get("id", ""))
            scls = " ".join(tag.get("class", [])) if isinstance(tag.get("class"), list) else str(tag.get("class", ""))
            if not (keyword_re.search(sid) or keyword_re.search(scls)):
                continue
            txt = (tag.get_text() or "").strip()
            if len(txt) < 120:
                continue
            if len(txt) > 48000:
                continue
            if self._looks_like_nav_text(txt[:900]):
                continue
            return tag

        blacklist = re.compile(r"header|footer|nav|menu|copyright|bread|top|bottom|sidebar|widget", re.I)
        candidates = soup.find_all(["div", "td", "article", "section"])
        best_tag, best_score = None, -1.0
        for tag in candidates:
            tag_id = str(tag.get("id", ""))
            tag_class = str(tag.get("class", ""))
            if blacklist.search(tag_id) or blacklist.search(tag_class):
                continue
            if self._tag_looks_like_page_chrome(tag):
                continue
            raw_txt = (tag.get_text() or "").strip()
            text_len = len(raw_txt)
            if text_len < 200 or text_len > 45000:
                continue
            q = self._score_content_quality(raw_txt[:2500])
            if q < 20:
                continue
            # 同等质量下略偏好「更像一篇文」的中等长度，避免选中整站外壳
            adj = q + min(text_len / 800.0, 18.0) - max(0.0, (text_len - 12000) / 9000.0)
            if adj > best_score:
                best_score, best_tag = adj, tag

        if best_tag is not None:
            return best_tag

        if full_mode:
            return None

        if soup.body and not self._tag_looks_like_page_chrome(soup.body):
            return soup.body
        return None

    def _is_decorative_image_content(self, pil_img):
        """根据图片内容分布过滤大面积背景/渐变/留白装饰图。"""
        try:
            img = pil_img.convert("RGB")
            w, h = img.size
            if w < 280 or h < 180:
                return True
            ratio = w / max(h, 1)
            if ratio > 3.4 or ratio < 0.28:
                return True

            # 抽样缩放后统计亮度/颜色方差；装饰底图通常“很亮 + 变化很小”
            sample = img.resize((64, 64))
            stat = ImageStat.Stat(sample)
            means = stat.mean  # R/G/B
            vars_ = stat.var
            brightness = sum(means) / 3.0
            avg_var = sum(vars_) / 3.0

            if brightness > 238 and avg_var < 220:
                return True

            # 白色像素占比过高，常见于底纹、留白页眉页脚
            px = sample.getdata()
            near_white = sum(1 for r, g, b in px if r >= 244 and g >= 244 and b >= 244)
            white_ratio = near_white / max(len(px), 1)
            if white_ratio > 0.86 and avg_var < 320:
                return True
        except Exception:
            # 内容判定失败不直接丢弃，交由其它规则决定
            return False
        return False

    def _collect_candidate_images(self, content_div, anchor_node, nearby_only: Optional[bool] = None):
        """收集正文区图片。

        若仅用 ``anchor_node.next_elements``，会漏掉「先配图、后段落」的常见版式（锚点在第一段文字上，
        前面的图永远遍历不到），表现为只能抓到一张或抓不到头图。

        ``nearby_only`` 为 None 时使用实例属性 ``nearby_images_only``；全量模式可由调用方传入 False 以遍历正文容器内全部 img。
        """
        if not content_div:
            return []
        all_in_div = content_div.find_all("img")
        use_nearby = self.nearby_images_only if nearby_only is None else bool(nearby_only)
        if not use_nearby or anchor_node is None:
            return all_in_div

        before: List = []
        scanned = 0
        for node in anchor_node.previous_elements:
            scanned += 1
            if scanned > self.nearby_scan_nodes:
                break
            if getattr(node, "name", None) != "img":
                continue
            if not self._img_under_content_root(node, content_div):
                continue
            before.append(node)
        before.reverse()

        after: List = []
        scanned = 0
        for node in anchor_node.next_elements:
            scanned += 1
            if scanned > self.nearby_scan_nodes:
                break
            if getattr(node, "name", None) != "img":
                continue
            if not self._img_under_content_root(node, content_div):
                continue
            after.append(node)

        merged = before + after
        seen: set[int] = set()
        ordered: List = []
        for im in merged:
            uid = id(im)
            if uid in seen:
                continue
            seen.add(uid)
            ordered.append(im)
        return ordered if ordered else all_in_div

    def _extract_wechat_text(self, soup, *, full_mode: bool = False):
        # 微信正文优先容器
        root = soup.find(id="js_content") or soup.find("div", class_=re.compile(r"rich_media_content", re.I))
        blocks = []
        if root:
            if full_mode:
                raw = root.get_text("\n", strip=True)
                full_t = self.prepare_full_report_text(raw)
                if full_t and self._score_content_quality(full_t) >= 28:
                    return full_t
                blocks, _ = self._collect_text_blocks(
                    root,
                    tags=["p", "section", "div", "span", "font"],
                    max_blocks=500,
                    max_chars=self.full_content_max_chars,
                )
            else:
                blocks, _ = self._collect_text_blocks(root, tags=["p", "section", "div"], max_blocks=10, max_chars=1600)
        text = self._dedup_sentences("\n".join(blocks))
        if self._score_content_quality(text) >= 35:
            return text
        if full_mode and text:
            return self.prepare_full_report_text(text)

        # 特效/反爬导致正文不可读时，降级到元数据，避免生成垃圾段落
        meta_text = self._extract_meta_text(soup)
        if meta_text:
            return meta_text if not full_mode else self.prepare_full_report_text(meta_text)
        return text

    def _extract_school_site_text(self, soup, *, full_mode: bool = False):
        main_el = soup.find("main")
        if main_el and not self._tag_looks_like_page_chrome(main_el):
            root = main_el
            inner_art = main_el.find("article")
            if inner_art and not self._tag_looks_like_page_chrome(inner_art):
                root = inner_art
            else:
                for sel in (".v_news_content", ".article-content", "#vsb_content", ".TRS_Editor", ".news_content"):
                    hit = main_el.select_one(sel)
                    if hit and not self._tag_looks_like_page_chrome(hit):
                        root = hit
                        break
            mb = 500 if full_mode else 12
            mc = self.full_content_max_chars if full_mode else 1800
            blocks, _ = self._collect_text_blocks(root, max_blocks=mb, max_chars=mc)
            text = self._dedup_sentences("\n".join(blocks))
            cleaned = self._strip_school_site_boilerplate(text)
            if cleaned and not self._looks_like_nav_text(cleaned):
                sc = self._score_content_quality(cleaned)
                sc += min(cleaned.count("。") * 4, 20)
                thresh = 30 if full_mode else 42
                if sc >= thresh:
                    return self.prepare_full_report_text(cleaned) if full_mode else cleaned
        selectors = [
            "#BodyLabel",
            ".n_list3",
            ".text",
            ".n_right",
            "#vsb_content",
            ".v_news_content",
            ".TRS_Editor",
            ".article-content",
            ".articleContent",
            ".news_content",
            ".content",
            "#content",
            "article",
        ]
        candidates = []
        for sel in selectors:
            for node in soup.select(sel):
                if self._tag_looks_like_page_chrome(node):
                    continue
                marker = " ".join(
                    [
                        str(node.get("id", "")),
                        " ".join(node.get("class", [])) if isinstance(node.get("class"), list) else str(node.get("class", "")),
                    ]
                ).lower()
                if re.search(r"nav|head|header|menu|footer|layout|sidebar|widget|bread", marker):
                    continue
                mb = 400 if full_mode else 10
                mc = self.full_content_max_chars if full_mode else 1500
                blocks, _ = self._collect_text_blocks(node, max_blocks=mb, max_chars=mc)
                text = self._dedup_sentences("\n".join(blocks))
                if not text:
                    continue
                cleaned = self._strip_school_site_boilerplate(text)
                if self._looks_like_nav_text(cleaned):
                    continue
                score = self._score_content_quality(cleaned)
                # 更偏向“有句号的连贯正文”
                score += min(cleaned.count("。") * 4, 20)
                candidates.append((score, len(cleaned), cleaned))
        if not candidates:
            return ""
        candidates.sort(key=lambda x: (x[0], x[1]), reverse=True)
        best = candidates[0][2]
        thresh = 28 if full_mode else 35
        if self._score_content_quality(best) < thresh:
            return ""
        return self.prepare_full_report_text(best) if full_mode else best

    def _looks_like_nav_text(self, text):
        t = self.normalize_article_text(text)
        if not t:
            return True
        nav_words = [
            "首页",
            "学院概况",
            "学院简介",
            "师资队伍",
            "学院动态",
            "通知公告",
            "学生工作",
            "组织架构",
            "成员介绍",
            "特色品牌",
            "实践活动",
            "传媒影音",
            "您的位置",
        ]
        nav_hits = sum(1 for w in nav_words if w in t)
        punct = len(re.findall(r"[。！？!?；;]", t))
        # 菜单/导航常是词串，句号极少且栏目词很多
        if nav_hits >= 4 and punct <= 1:
            return True
        # 过多短词堆叠也视作导航噪音
        chunks = [x for x in re.split(r"\s+", t) if x]
        short_chunks = sum(1 for x in chunks if len(x) <= 6)
        if len(chunks) >= 20 and short_chunks / max(len(chunks), 1) >= 0.8 and punct <= 1:
            return True
        return False

    def _strip_school_site_boilerplate(self, text):
        t = self.normalize_article_text(text)
        if not t:
            return t
        # 去掉学校站点底部联系方式和备案类信息
        t = re.sub(r"电话：.*", "", t)
        t = re.sub(r"招生录取咨询专线：.*", "", t)
        t = re.sub(r"邮箱：.*", "", t)
        t = re.sub(r"邮 编：.*", "", t)
        t = re.sub(r"地 址：.*", "", t)
        t = re.sub(r"©\s*\d{4}.*", "", t)
        t = re.sub(r"Produced By.*", "", t, flags=re.I)
        t = re.sub(r"publishdate[:：].*", "", t, flags=re.I)
        return self.normalize_article_text(t)

    def summarize_for_briefing(self, text, max_chars: Optional[int] = None):
        cap = self.summary_max_chars if max_chars is None else int(max_chars)
        cap = max(80, cap)
        t = self._dedup_sentences(self.normalize_article_text(text))
        if not t:
            return t
        # 尝试切句，保留前几句，避免把整篇公众号全文塞进简报。
        parts = [s.strip() for s in re.split(r"(?<=[。！？!?])", t) if s.strip()]
        if not parts:
            return t[:cap]
        out, total = [], 0
        for s in parts:
            if total + len(s) > cap and out:
                break
            out.append(s)
            total += len(s)
            if len(out) >= 5:
                break
        res = "".join(out).strip()
        if not res:
            res = t[:cap].strip()
        return res

    def prepare_full_report_text(self, text: str) -> str:
        """全量模式：去重句、去底部样板废话后截断到上限，不做「摘要五句」裁剪。"""
        t = self._dedup_sentences(self.normalize_article_text(text))
        if not t:
            return ""
        t = self._strip_school_site_boilerplate(t)
        cap = self.full_content_max_chars
        if len(t) > cap:
            t = t[:cap].rstrip() + "\n\n（正文已按上限截断，余下内容请打开原链接查看。）"
        return t

    def is_junk_paragraph(self, text):
        text = self.normalize_article_text(text)
        if self._looks_like_nav_text(text):
            return True
        junk_keywords = [
            "首页 >",
            "首页>",
            "当前位置",
            "您现在的位置",
            "来源：",
            "作者：",
            "编辑：",
            "浏览：",
            "发布时间：",
            "打印",
            "关闭",
            "分享到：",
            "一审：",
            "二审：",
            "三审：",
            "审核：",
            "编辑：",
            "责编：",
            "来源：",
            "作者：",
            "融媒体中心",
            "上一篇",
            "下一篇",
            "阅读原文",
            "点击蓝字",
            "返回搜狐",
            "写留言",
            "写评论",
            "赞在看",
            "在看",
            "轻点两下",
            "预览时标签不可点",
            "收录于话题",
        ]
        for kw in junk_keywords:
            if kw in text:
                return True
        if len(text) < 18:
            return True
        valid_endings = ("。", "！", "？", '"', "…", ".")
        if len(text) > 50 and not text.endswith(valid_endings):
            if re.search(r"\d{4}-\d{2}-\d{2}", text):
                return True
        return False

    def get_article_content(self, url):
        self.article_error = ''
        if os.environ.get('SEQARA_OFFLINE') == '1':
            self.article_error = '离线模式下请在表格的“活动内容”或“新闻正文”列填写正文。'
            return '', []
        try:
            full = self.full_content_mode
            parsed = urlparse(url)
            referer = f"{parsed.scheme}://{parsed.netloc}/" if parsed.netloc else url
            response = self.session.get(
                url,
                timeout=self.fetch_timeout,
                verify=False,
                headers={"Referer": referer},
            )
            response.raise_for_status()
            html_text = self._decode_response_html(response)
            soup = BeautifulSoup(html_text, "html.parser")
            host = (urlparse(url).netloc or "").lower()
            raw_page_text = self.normalize_article_text(soup.get_text(" ", strip=True))
            content_div = None

            if "mp.weixin.qq.com" in host:
                wx_text = self._extract_wechat_text(soup, full_mode=full)
                if "环境异常" in raw_page_text and "去验证" in raw_page_text and self._score_content_quality(wx_text) < 35:
                    self.article_error = '微信页面触发验证，未获取正文。'
                    return '', []
                # 全量模式禁止用整页文本兜底，否则会混入导航、侧栏、页脚等噪音。
                if full:
                    content_text = wx_text or ""
                else:
                    content_text = self.summarize_for_briefing(wx_text or raw_page_text)
                if self._looks_broken_wechat_text(content_text):
                    self.article_error = '微信页面正文识别失败，请补充本地正文。'
                    return '', []
                content_div = soup.find(id="js_content") or soup.find("div", class_=re.compile(r"rich_media_content", re.I))
            else:
                school_text = self._extract_school_site_text(soup, full_mode=full)
                if school_text:
                    content_text = school_text if full else self.summarize_for_briefing(school_text)
                else:
                    content_text = ""

            jd_text = ""
            if "mp.weixin.qq.com" not in host:
                jd_text = self._extract_json_ld_article_text(soup)
            if jd_text:
                if full:
                    jd_use = self.prepare_full_report_text(jd_text)
                else:
                    jd_use = self.summarize_for_briefing(jd_text)
                if jd_use and self._score_content_quality(jd_use) > self._score_content_quality(content_text or ""):
                    content_text = jd_use

            if not content_div:
                content_div = self._resolve_content_div(soup, host, full_mode=full)

            anchor_node = None
            if not content_text and content_div:
                mb = 320 if full else 6
                mc = self.full_content_max_chars if full else 900
                paragraph_texts, anchor_node = self._collect_text_blocks(content_div, max_blocks=mb, max_chars=mc)
                if paragraph_texts:
                    joined = "\n".join(paragraph_texts)
                    content_text = self.prepare_full_report_text(joined) if full else self.summarize_for_briefing(joined)

            if not content_text or self._score_content_quality(content_text) < 25:
                meta_text = self._extract_meta_text(soup)
                if meta_text:
                    if full:
                        content_text = self.prepare_full_report_text(meta_text)
                    else:
                        content_text = self.summarize_for_briefing(meta_text, max_chars=min(260, self.summary_max_chars))

            if not content_text:
                self.article_error = '未能抓取有效正文，请补充本地正文。'
                return '', []

            images_data = []
            img_cap = self.full_mode_max_images if full else self.max_images_per_article
            nearby_eff = (not full) and self.nearby_images_only
            if content_div:
                candidate_imgs = self._collect_candidate_images(content_div, anchor_node, nearby_only=nearby_eff)
                count = 0
                for img in candidate_imgs:
                    if count >= img_cap:
                        break
                    src = self._img_url_from_tag(img)
                    if not src:
                        continue
                    full_src = urljoin(url, src)
                    if self._is_decorative_img_tag(img, full_src):
                        continue
                    ext = os.path.splitext(urlparse(full_src).path)[1].lower()
                    if ext == ".gif":
                        continue
                    try:
                        img_resp = self.session.get(
                            full_src,
                            timeout=max(8, self.fetch_timeout - 4),
                            verify=False,
                            headers={"Referer": url},
                        )
                        if img_resp.status_code != 200:
                            continue
                        ctype = (img_resp.headers.get("Content-Type") or "").lower()
                        raw = img_resp.content
                        if not raw:
                            continue
                        # 部分 CDN 返回 application/octet-stream，仍以 PIL 能否打开为准
                        if "image" not in ctype:
                            try:
                                Image.open(BytesIO(raw)).verify()
                            except Exception:
                                continue
                        img_bytes = BytesIO(raw)
                        try:
                            pil_img = Image.open(img_bytes)
                            pil_img.load()
                            w, h = pil_img.size
                            if w >= 280 and h >= 180:
                                if self._is_decorative_image_content(pil_img):
                                    continue
                                img_bytes.seek(0)
                                images_data.append(img_bytes)
                                count += 1
                        except Exception:
                            continue
                    except Exception:
                        continue
            return content_text, images_data
        except Exception as e:
            self.log(f"  ⚠️ 抓取异常: {e}")
            self.article_error = f'链接访问失败：{e}'
            return '', []

    @staticmethod
    def _safe_filename(name: str, fallback: str = "未命名活动") -> str:
        return safe_file_stem(name, fallback)

    @staticmethod
    def _activity_name(row):
        value = row.get('活动名称')
        return str(value).strip() if pd.notna(value) and str(value).strip() else '未命名活动'

    @staticmethod
    def _unique_filepath(folder: str, filename: str) -> str:
        path = os.path.join(folder, filename)
        if not os.path.exists(path):
            return path
        stem, ext = os.path.splitext(filename)
        for i in range(2, 1000):
            candidate = os.path.join(folder, f"{stem}_{i}{ext}")
            if not os.path.exists(candidate):
                return candidate
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return os.path.join(folder, f"{stem}_{stamp}{ext}")

    def _resolve_header_school_name(self, df: pd.DataFrame) -> str:
        header_school_name = ""
        if "学校名称" in df.columns:
            valid_schools = df["学校名称"].dropna()
            if not valid_schools.empty:
                val = str(valid_schools.iloc[0]).strip()
                if val and val.lower() != "nan":
                    header_school_name = val
        elif "学院名称" in df.columns:
            valid_colleges = df["学院名称"].dropna()
            if not valid_colleges.empty:
                val = str(valid_colleges.iloc[0]).strip()
                if val and val.lower() != "nan":
                    header_school_name = val
        return header_school_name

    @staticmethod
    def _cell_text(value):
        return '' if value is None or pd.isna(value) else str(value).strip()

    def _row_school(self, row):
        return self._cell_text(row.get('学校名称')) or self._cell_text(row.get('学院名称'))

    def _merged_doc_title(self, header_school_name: str) -> str:
        if self.full_content_mode:
            return f"{header_school_name}示例设计产业学院活动简报（全量）"
        return f"{header_school_name}示例设计产业学院活动简报"

    def _activity_doc_title(self, header_school_name: str, activity_name: str) -> str:
        base = f"{header_school_name}示例设计产业学院 - {activity_name}"
        if self.full_content_mode:
            return f"{base}（全量）"
        return base

    def _new_document(self, doc_title: str) -> Document:
        doc = Document()
        title_para = doc.add_paragraph()
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_para.add_run(doc_title)
        self.set_font(title_run, "仿宋", size=22, bold=True)
        return doc

    def _append_activity_section(
        self,
        doc: Document,
        row: Any,
        *,
        index: int,
        total_rows: int,
        section_number: int,
        numbered_subtitle: bool,
    ) -> bool:
        raw_link_text = row.get("新闻链接", "")
        link, status = self.extract_url(raw_link_text)
        activity_name = self._activity_name(row)
        content_text = next((self._cell_text(row.get(k)) for k in
            ('活动内容', '新闻正文', '正文', '活动总结', '内容') if self._cell_text(row.get(k))), '')
        images = []
        if not content_text and link:
            content_text, images = self.get_article_content(link)
        if not content_text:
            reason = getattr(self, 'article_error', '') or status or '缺少正文'
            self.result['failed'].append(f'第 {index + getattr(self, "_source_header_row", 1) + 1} 行 [{activity_name}]：{reason}')
            self.log(f"⚠️ 跳过第 {index + 2} 行 [{activity_name}]: {reason}")
            return False
        raw_time = row.get("活动时间", "")
        activity_time = self.smart_clean_time(raw_time)
        self.log(f"正在处理 ({index + 1}/{total_rows}): {activity_name}")

        def get_count_str(val, label):
            num = parse_headcount(val)
            return f"{label}：{num}人" if num > 0 else ""

        try:
            teacher_count = get_count_str(row.get("服务教师（人数）", 0), "服务教师")
            student_count = get_count_str(row.get("服务学生（人数）", 0), "服务学生")
        except ValueError as exc:
            reason = f'第 {index + self._source_header_row + 1} 行 [{activity_name}]：{exc}'
            self.result['failed'].append(reason)
            self.log(reason)
            return False
        doc.add_paragraph()
        subtitle_para = doc.add_paragraph()
        school = self._row_school(row)
        activity_label = f'{school} · {activity_name}' if getattr(self, 'multiple_schools', False) else activity_name
        subtitle_text = f"{section_number}. {activity_label}" if numbered_subtitle else activity_label
        subtitle_run = subtitle_para.add_run(subtitle_text)
        self.set_font(subtitle_run, "仿宋", size=16, bold=True)
        if self.keep_activity_meta:
            info_text = "  ".join(f"{activity_time}  {teacher_count}  {student_count}".strip().split())
            info_para = doc.add_paragraph()
            info_para.paragraph_format.first_line_indent = Pt(24)
            self.set_font(info_para.add_run(info_text), "仿宋", size=12, bold=False)
        content_para = doc.add_paragraph()
        content_para.paragraph_format.first_line_indent = Pt(24)
        self.set_font(content_para.add_run(content_text), "仿宋", size=12, bold=False)
        if self.keep_activity_meta and link:
            link_para = doc.add_paragraph()
            link_para.paragraph_format.first_line_indent = Pt(24)
            self.set_font(link_para.add_run(f"新闻链接：{link}"), "仿宋", size=12, bold=False)
        if images:
            img_para = doc.add_paragraph()
            img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for img_stream in images:
                try:
                    img_stream.seek(0)
                    run = img_para.add_run()
                    run.add_picture(img_stream, width=Inches(5))
                    run.add_text("\n")
                except Exception as e:
                    self.log(f"  图片插入失败: {e}")
        if self.row_delay_sec > 0:
            time.sleep(self.row_delay_sec)
        return True

    def _generate_merged(self, df: pd.DataFrame, output_folder: str, header_school_name: str) -> None:
        safe_name = self._safe_filename(header_school_name, "学校")
        if self.full_content_mode:
            school_name_for_file = f"{safe_name}_活动全量简报"
        else:
            school_name_for_file = f"{safe_name}_活动简报"
        doc = self._new_document(self._merged_doc_title(header_school_name))
        valid_count = 0
        total_rows = len(df)
        for index, row in df.iterrows():
            if self.stop_flag:
                self.log("任务已停止")
                return
            if self._append_activity_section(
                doc,
                row,
                index=index,
                total_rows=total_rows,
                section_number=valid_count + 1,
                numbered_subtitle=True,
            ):
                valid_count += 1
        if self.stop_flag:
            self.log('任务已停止，本次合并文件未保存。')
            return
        if valid_count == 0:
            self.log("❌ 未生成简报：所有活动均缺少可用正文。请补充活动内容或检查新闻链接。")
            self._dialog("warning", "生成失败", "未生成简报：所有活动均缺少可用正文，请查看失败明细。")
            return
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_filename = f"{school_name_for_file}_{timestamp}.docx"
        output_path = self._unique_filepath(output_folder, output_filename)
        try:
            with staged_outputs(output_path) as (staged,):
                doc.save(staged)
            self.result['success'] = valid_count
            self.result['artifacts'].append(output_path)
            outcome = '部分成功' if self.result['failed'] else '成功'
            self.log(f"\n{outcome}：已生成 {valid_count}/{total_rows} 条活动，文件已保存至:\n{output_path}")
            kind_label = "全量简报" if self.full_content_mode else "简报"
            self._dialog("info", outcome, f"{kind_label}已生成 {valid_count}/{total_rows} 条活动。\n文件名：{Path(output_path).name}\n失败明细见日志。" if self.result['failed'] else f"{kind_label}已生成！\n文件名：{Path(output_path).name}")
        except Exception as e:
            self.result['failed'].append(f'合并简报保存失败：{e}')
            self.log(f"❌ 保存文件失败: {e}")
            self._dialog("error", "错误", f"保存失败: {e}")

    def _generate_per_activity(self, df: pd.DataFrame, output_folder: str, header_school_name: str) -> None:
        saved_files: List[str] = []
        total_rows = len(df)
        for index, row in df.iterrows():
            if self.stop_flag:
                self.log("任务已停止")
                return
            activity_name = self._activity_name(row)
            doc_title = self._activity_doc_title(self._row_school(row), activity_name)
            doc = self._new_document(doc_title)
            if not self._append_activity_section(
                doc,
                row,
                index=index,
                total_rows=total_rows,
                section_number=1,
                numbered_subtitle=False,
            ):
                continue
            if self.stop_flag:
                self.log('任务已停止，当前活动文件未保存。')
                return
            output_filename = f"{self._safe_filename(activity_name)}.docx"
            output_path = self._unique_filepath(output_folder, output_filename)
            try:
                with staged_outputs(output_path) as (staged,):
                    doc.save(staged)
                saved_files.append(output_path)
                self.result['success'] += 1
                self.result['artifacts'].append(output_path)
                self.log(f"  ✅ 已保存: {os.path.basename(output_path)}")
            except Exception as e:
                self.result['failed'].append(f'{activity_name}：保存失败：{e}')
                self.log(f"  ❌ [{activity_name}] 保存失败: {e}")
        if not saved_files:
            self.log("❌ 未生成简报，请补充可用正文或查看保存失败明细。")
            self._dialog("warning", "生成失败", "未生成简报，请查看失败明细。")
            return
        outcome = '部分成功' if self.result['failed'] else '成功'
        self.log(f"\n{outcome}：生成 {len(saved_files)}/{total_rows} 个 Word 文件，保存在:\n{output_folder}")
        preview = "\n".join(os.path.basename(p) for p in saved_files[:8])
        if len(saved_files) > 8:
            preview += f"\n…等共 {len(saved_files)} 个文件"
        kind_label = "全量简报" if self.full_content_mode else "简报"
        self._dialog("info", outcome, f"已按活动生成 {len(saved_files)}/{total_rows} 个{kind_label}。\n\n{preview}")

    def generate(self, excel_path, output_folder):
        self.result = {'success': 0, 'failed': [], 'artifacts': [], 'status': 'failed'}
        self.article_error = ''
        excel_path = str(excel_path)
        self._source_header_row = 1
        self.log(f"正在读取文件: {excel_path} ...")
        try:
            if Path(excel_path).suffix.lower() == '.csv':
                df = pd.read_csv(excel_path, dtype=str)
            else:
                try:
                    df = pd.read_excel(excel_path, header=1, engine="openpyxl", dtype=str)
                    if "活动名称" not in df.columns or not any(k in df.columns for k in ('新闻链接', '活动内容', '新闻正文', '正文', '活动总结', '内容')):
                        raise ValueError("Header mismatch")
                    self._source_header_row = 2
                except Exception:
                    self.log("提示：尝试标准格式读取...")
                    df = pd.read_excel(excel_path, header=0, engine="openpyxl", dtype=str)
        except Exception as e:
            self.log(f"❌ 读取失败: {e}")
            self.result['failed'].append(str(e))
            return self.result
        if not any(k in df.columns for k in ('新闻链接', '活动内容', '新闻正文', '正文', '活动总结', '内容')):
            self.result['failed'].append('请提供新闻链接或活动内容／新闻正文列。')
            self.log('❌ ' + self.result['failed'][-1])
            return self.result
        missing = [str(i + self._source_header_row + 1) for i, row in df.iterrows() if not self._row_school(row)]
        if df.empty or missing:
            message = '没有活动记录。' if df.empty else '请补全学校名称或学院名称，第 ' + '、'.join(missing) + ' 行。'
            self.result['failed'].append(message)
            self.log('❌ ' + message)
            self._dialog('warning', '无法生成简报', message)
            return self.result
        schools = list(dict.fromkeys(self._row_school(row) for _, row in df.iterrows()))
        self.multiple_schools = len(schools) > 1
        Path(output_folder).mkdir(parents=True, exist_ok=True)
        if self.full_content_mode:
            self.log(
                "📎 已启用「全量报告」：在识别出的正文容器内尽量保留全文并取配图（不合并整页导航/侧栏）；"
                "文档更大、耗时更长，如遇大页可将「单页超时」适当调高。"
            )
        if self.split_by_activity:
            self.log("📁 已启用「按活动分文件」：每个活动单独保存为一个 Word，文件名与活动名称一致。")
        header_school_name = '多院校' if self.multiple_schools else schools[0]
        if self.split_by_activity:
            self._generate_per_activity(df, output_folder, header_school_name)
        else:
            self._generate_merged(df, output_folder, header_school_name)
        self.result['status'] = 'cancelled' if self.stop_flag else (('partial' if self.result['failed'] else 'completed') if self.result['success'] else 'failed')
        return self.result
