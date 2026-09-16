# -*- coding: utf-8 -*-
"""
与「工具集V4.3 CTK替换新UI底层.py」中同名逻辑保持同步的可导入业务层。
供 PySide6 逐页重构版调用；避免 import CTK 主文件（会初始化 Tk 界面）。
"""
from __future__ import annotations

import os
import re
from pathlib import Path
from collections import defaultdict
from datetime import datetime

import pandas as pd

from nv_school_utils import school_from_department
from nv_activity_types import normalize_activity_types
from nv_data_safety import new_output_path, parse_headcount, staged_outputs
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

try:
    from rapidfuzz import fuzz, process
except ImportError as exc:
    raise RuntimeError("缺少 rapidfuzz；请在构建或部署阶段安装完整依赖。") from exc

LOG_DATA = "error_log_data.txt"
INFO_LOG = "info_log.txt"

post_ui = None  # Callable[[Callable[[], None]], None]
show_info = None  # (title, message) -> None
show_error = None  # (title, message) -> None


def _notify_status(status_label, *, text=None, text_color=None) -> None:
    """PySide6 下后台线程勿依赖 post_ui(lambda…)；带 _thread_safe_configure 的适配器会自行排队到主线程。"""
    if status_label is None:
        return
    if getattr(status_label, "_thread_safe_configure", False):
        status_label.configure(text=text, text_color=text_color)
        return
    if post_ui:
        post_ui(lambda sl=status_label, t=text, c=text_color: sl.configure(text=t, text_color=c))


def log_data(msg: str) -> None:
    with open(LOG_DATA, "a", encoding="utf-8") as f:
        f.write(f"[{datetime.now()}] {msg}\n")


def log_info(msg: str) -> None:
    with open(INFO_LOG, "a", encoding="utf-8") as f:
        f.write(f"[{datetime.now()}] {msg}\n")


def _clean_output_basename(stem, default="示例设计产业学院活动整理", max_len=120):
    s = (stem or "").strip() or default
    for ch in r'\/:*?"<>|':
        s = s.replace(ch, "")
    s = s.strip(". ")
    return (s or default)[:max_len]


def _stat_safe_int(val, default):
    try:
        return int(float(str(val).strip()))
    except (TypeError, ValueError):
        return default


def _stat_safe_float(val, default):
    try:
        return float(str(val).strip())
    except (TypeError, ValueError):
        return default


def _stat_sanitize_hex(color, default):
    if not color:
        return default
    s = str(color).strip().lstrip("#")
    if len(s) == 6 and all(c in "0123456789ABCDEFabcdef" for c in s):
        return s.upper()
    return default


_DEFAULT_STAT_ACTIVITY_TYPES = [
    "请进来", "走出去", "研学", "展览", "企业研学", "师资培训", "训练营", "工作坊",
    "产业学院赛事", "其他竞赛", "论坛", "国际交流", "人培研讨会",
    "产教融合课程", "校企合作", "成果展", "企业挂职",
    "协同育人", "双师赋能", "设计工坊", "校企课程", "产业赛事",
    "学科竞赛", "成果转化", "国际合作", "校企共建", "实验室建设", "其他",
    "实训营", "课程", "赛事", "品牌活动","人培研讨",
]


def format_date(date_str):
    try:
        if isinstance(date_str, str):
            parts = date_str.strip().split("/")
            if len(parts) == 3:
                year, month, day = parts
                return f"{int(year)}年{int(month)}月{int(day)}日"
        return date_str
    except Exception:
        return date_str


def parse_datetime(x):
    try:
        return pd.to_datetime(x, format="%Y/%m/%d")
    except Exception:
        return pd.NaT


def style_excel(
    path,
    *,
    header_font_name="微软雅黑",
    header_font_size=11,
    header_font_color="FFFFFF",
    body_font_name="微软雅黑",
    body_font_size=10,
    header_fill_hex="4F81BD",
    body_row_height=38,
):
    def _rgb6(val, default):
        if not val:
            return default
        s = str(val).strip().lstrip("#").upper()
        return s if len(s) == 6 and all(c in "0123456789ABCDEF" for c in s) else default

    hb = _rgb6(header_fill_hex, "4F81BD")
    hc = _rgb6(header_font_color, "FFFFFF")
    hname = (header_font_name or "微软雅黑")[:31]
    bname = (body_font_name or "微软雅黑")[:31]
    try:
        hsz = int(header_font_size)
    except (TypeError, ValueError):
        hsz = 11
    try:
        bsz = int(body_font_size)
    except (TypeError, ValueError):
        bsz = 10
    try:
        brh = float(body_row_height)
    except (TypeError, ValueError):
        brh = 38
    if brh < 10:
        brh = 38

    wb = openpyxl.load_workbook(path)
    ws = wb.active
    header_font = Font(name=hname, bold=True, color=hc, size=hsz)
    body_font = Font(name=bname, size=bsz)
    header_fill = PatternFill(start_color=hb, end_color=hb, fill_type="solid")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border = Border(*(Side(style="thin", color="000000"),) * 4)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = border
    time_col_idx = None
    for idx, cell in enumerate(ws[1], start=1):
        if cell.value == "活动时间":
            time_col_idx = idx
            break
    month_colors = {
        1: "E2F0D9", 2: "FFF2CC", 3: "FBE5D6", 4: "E6E6FA",
        5: "9DC3E6", 6: "F4B183", 7: "E0EEE0", 8: "FFEFD5",
        9: "F0F8FF", 10: "F5E6E8", 11: "FFF0F5", 12: "F8F8FF",
    }
    default_color = "FFFFFF"
    for col in ws.columns:
        col_name = col[0].value
        col_letter = get_column_letter(col[0].column)
        if col_name == "活动时间":
            ws.column_dimensions[col_letter].width = 27
        elif col_name == "新闻链接":
            ws.column_dimensions[col_letter].width = 60
        elif col_name == "服务学生（人数）":
            ws.column_dimensions[col_letter].width = 10
        elif col_name in ["学校名称", "学院名称", "活动类型"]:
            ws.column_dimensions[col_letter].width = 22
        elif col_name in ["活动性质", "覆盖专业"]:
            ws.column_dimensions[col_letter].width = 15
        else:
            lens = [len(str(cell.value)) for cell in col if cell.value]
            max_len = max(lens) if lens else 8
            ws.column_dimensions[col_letter].width = max_len + 2
        for i, cell in enumerate(col, start=1):
            cell.alignment = align_center
            cell.font = body_font
            cell.border = border
            ws.row_dimensions[i].height = brh
            if time_col_idx and col[0].column == time_col_idx and i > 1 and cell.value:
                try:
                    if "/" in str(cell.value):
                        y, m, d = [int(x) for x in str(cell.value).split("/")]
                        cell.value = f"{y}年{m}月{d}日"
                except Exception:
                    pass
                try:
                    month = int(str(cell.value).split("年")[1].split("月")[0])
                except Exception:
                    month = 0
                color = month_colors.get(month, default_color)
                cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    wb.save(path)


def extract_school(department):
    return school_from_department(department)


def clean_approval_no(x):
    if pd.isna(x):
        return ""
    x = str(x).strip().replace(".0", "")
    x = re.sub(r"[^\d]", "", x)
    return x if len(x) >= 6 else ""


def _read_activity_excel(path):
    """Expand only actual merged cells, never neighbouring independent records."""
    frame = pd.read_excel(path)
    if str(path).lower().endswith('.xlsx'):
        book = openpyxl.load_workbook(path, data_only=True)
        try:
            sheet = book.worksheets[0]
            for area in sheet.merged_cells.ranges:
                value = sheet.cell(area.min_row, area.min_col).value
                for row in range(max(2, area.min_row), area.max_row + 1):
                    for col in range(area.min_col, area.max_col + 1):
                        if row - 2 < len(frame) and col - 1 < len(frame.columns):
                            frame.iloc[row - 2, col - 1] = value
        finally:
            book.close()
    return frame


def match_and_clean(
    app_file,
    reim_file,
    output_dir,
    status_label,
    *,
    fuzzy_threshold=90,
    output_basename="示例设计产业学院活动整理",
    style_excel_row_height=38,
    style_header_fill_hex="4F81BD",
    style_header_font_color="FFFFFF",
    style_header_font_size=11,
    style_body_font_size=10,
):
    try:
        ft = _stat_safe_int(fuzzy_threshold, 90)
        ft = max(50, min(100, ft))
        stem = _clean_output_basename(output_basename)
        df1 = _read_activity_excel(app_file)
        df2 = _read_activity_excel(reim_file)
        required_cols_1 = ["活动名称", "学院名称", "活动类型", "开始时间", "服务教师（人数）", "服务学生（人数）",
                           "活动总额", "申请人部门", "审批编号"]
        required_cols_2 = ["活动名称", "新闻链接", "关联申请单"]
        optional_cols = ["活动性质", "覆盖专业"]
        for col in required_cols_1:
            if col not in df1.columns:
                raise ValueError(f"活动申请表缺少字段：{col}")
        for col in required_cols_2:
            if col not in df2.columns:
                raise ValueError(f"活动报销表缺少字段：{col}")
        for col in required_cols_1 + ["申请人部门"] + optional_cols:
            if col in df1.columns:
                df1[col] = df1[col].replace(r"^\s*$", pd.NA, regex=True).fillna("")
            elif col in optional_cols:
                df1[col] = ""
        for col in required_cols_2:
            if col in df2.columns:
                df2[col] = df2[col].replace(r"^\s*$", pd.NA, regex=True).fillna("")
        df1["学校名称"] = df1["申请人部门"].apply(extract_school).fillna("")
        df1["审批编号"] = df1["审批编号"].apply(clean_approval_no)
        df2['匹配审批编号'] = df2['关联申请单'].apply(
            lambda value: list(dict.fromkeys(re.findall(r'\d{6,}', str(value)))) if pd.notna(value) else [])
        df2 = df2.explode('匹配审批编号').fillna({'匹配审批编号': ''})
        df2["学校名称"] = df2["学校名称"] if "学校名称" in df2.columns else ""
        missing_ids = df1.index[df1['审批编号'] == ''].tolist()
        if missing_ids:
            raise ValueError('申请表第 ' + '、'.join(str(i + 2) for i in missing_ids[:10]) + ' 行审批编号缺失或无效，请补全后重试。')
        df1 = df1.drop_duplicates(subset='审批编号', keep='first')
        valid_news = df2['新闻链接'].apply(lambda x: isinstance(x, str) and bool(re.match(r'^https?://', x.strip(), re.I)))
        # Keep unlinked reimbursements for the explicit name+school fallback.
        df2 = df2.loc[valid_news].copy()
        df2['新闻链接'] = df2['新闻链接'].str.strip()
        linked = df2[df2['匹配审批编号'] != ''].drop_duplicates(subset='匹配审批编号', keep='first')
        merged = pd.merge(df1, linked[['匹配审批编号', '新闻链接']], left_on='审批编号', right_on='匹配审批编号', how='left')
        # iterrows 的 idx 是行标签；lock 是按位置的列表。df1 筛选后索引常不连续，用标签当下标会错位或 IndexError。
        merged = merged.reset_index(drop=True)
        lock = merged["新闻链接"].apply(lambda x: isinstance(x, str) and x.startswith("http")).tolist()
        merged["匹配方式"] = merged["新闻链接"].apply(
            lambda x: "审批编号匹配" if pd.notna(x) and str(x).startswith("http") else "未匹配")
        df2_tri = df2[["活动名称", "新闻链接", "学校名称"]].copy()
        df2_tri["活动名称_clean"] = df2_tri["活动名称"].astype(str).str.lower().str.replace(" ", "")
        # 旧实现：对每个申请行反复全表扫描 + iterrows + Python 层 fuzz.ratio，数据稍大即极慢。
        by_school: dict[str, list[tuple[str, object]]] = defaultdict(list)
        all_clean_url: list[tuple[str, object]] = []
        exact_map: dict[tuple[str, str], object] = {}
        for cl, url, sc in zip(
            df2_tri["活动名称_clean"],
            df2_tri["新闻链接"],
            df2_tri["学校名称"].astype(str),
        ):
            cl_s, sc_s = str(cl), str(sc)
            pair = (cl_s, url)
            by_school[sc_s].append(pair)
            all_clean_url.append(pair)
            ek = (sc_s, cl_s)
            if ek not in exact_map:
                exact_map[ek] = url

        n = len(merged)
        act_list = merged["活动名称"].astype(str).str.lower().str.replace(" ", "").tolist()
        school_list = merged["学校名称"].astype(str).tolist()
        news_links = merged["新闻链接"].tolist()
        match_way = merged["匹配方式"].tolist()

        for idx in range(n):
            if lock[idx]:
                continue
            u = exact_map.get((school_list[idx], act_list[idx]))
            if u is not None:
                news_links[idx] = u
                match_way[idx] = "活动名称精准+学校匹配"
                lock[idx] = True

        # 按 (学校, 清洗后活动名) 分组：相同键的多行结果一致，避免对大表逐行重复扫描/模糊计算（大数据下否则极慢）
        _grp: dict[tuple[str, str], list[int]] = defaultdict(list)
        for idx in range(n):
            if lock[idx]:
                continue
            act_name = act_list[idx]
            school = school_list[idx]
            if not act_name or act_name == "nan":
                continue
            _grp[(school, act_name)].append(idx)

        for (school, act_name), idxs in _grp.items():
            if lock[idxs[0]]:
                continue
            cands = by_school.get(school, [])
            for cl, url in cands:
                if act_name in cl:
                    for idx in idxs:
                        news_links[idx] = url
                        match_way[idx] = "活动名称包含+学校匹配"
                        lock[idx] = True
                    break

        for (school, act_name), idxs in _grp.items():
            if lock[idxs[0]]:
                continue
            cands = all_clean_url if not school else by_school.get(school, [])
            if not cands:
                continue
            names = [p[0] for p in cands]
            hit = process.extractOne(act_name, names, scorer=fuzz.ratio, score_cutoff=float(ft))
            if hit is None:
                continue
            _m, score, j = hit
            url = cands[j][1]
            if url:
                for idx in idxs:
                    news_links[idx] = url
                    match_way[idx] = f"模糊匹配({int(score)})"
                    lock[idx] = True

        merged["新闻链接"] = news_links
        merged["匹配方式"] = match_way

        def safe_event_time(x):
            from nv_dashboard_core import _parse_date
            parsed = _parse_date(x)
            return parsed.isoformat() if parsed else ''

        cleaned = pd.DataFrame()
        cleaned['审批编号'] = merged['审批编号']
        cleaned["序号"] = range(1, len(merged) + 1)
        cleaned["活动名称"] = merged["活动名称"]
        cleaned["学校名称"] = merged["学校名称"]
        cleaned["学院名称"] = merged["学院名称"]
        cleaned["活动类型"] = merged["活动类型"].apply(normalize_activity_types)
        cleaned["活动性质"] = merged["活动性质"] if "活动性质" in merged.columns else ""
        cleaned["覆盖专业"] = merged["覆盖专业"] if "覆盖专业" in merged.columns else ""
        cleaned["活动时间"] = merged["开始时间"].apply(safe_event_time)
        cleaned["新闻链接"] = merged["新闻链接"]
        cleaned["服务教师（人数）"] = merged["服务教师（人数）"]
        cleaned["服务学生（人数）"] = merged["服务学生（人数）"]
        cleaned["费用"] = merged["活动总额"]
        cleaned["匹配方式"] = merged["匹配方式"]
        cleaned["活动时间_原始"] = cleaned["活动时间"].apply(parse_datetime)
        cleaned = cleaned.drop_duplicates(subset=['审批编号'], keep='first')
        cleaned = cleaned.sort_values(by=["活动时间_原始", "活动名称"]).reset_index(drop=True)
        cleaned["序号"] = range(1, len(cleaned) + 1)
        cleaned["活动时间"] = cleaned["活动时间"].apply(format_date)
        cleaned = cleaned.drop(columns=["活动时间_原始"])
        if not os.path.exists(output_dir):
            os.makedirs(output_dir, exist_ok=True)
        out_xlsx = new_output_path(output_dir, f'{stem}.xlsx', (app_file, reim_file))
        out_csv = new_output_path(output_dir, f'{out_xlsx.stem}.csv', (app_file, reim_file))
        with staged_outputs(out_xlsx, out_csv) as (staged_xlsx, staged_csv):
            cleaned.to_excel(staged_xlsx, index=False)
            cleaned.to_csv(staged_csv, index=False, encoding="utf-8-sig")
            _srh = _stat_safe_int(style_excel_row_height, 38)
            _srh = max(10, min(200, _srh))
            style_excel(
                staged_xlsx,
                header_fill_hex=style_header_fill_hex,
                header_font_color=style_header_font_color,
                header_font_size=_stat_safe_int(style_header_font_size, 11),
                body_font_size=_stat_safe_int(style_body_font_size, 10),
                body_row_height=_srh,
            )
        success_cnt = (cleaned["匹配方式"] != "未匹配").sum()
        fail_cnt = (cleaned["匹配方式"] == "未匹配").sum()
        _msg_ok = f"✅ 整理完成：{out_xlsx.name}（模糊≥{ft}）成功{success_cnt}，未匹配{fail_cnt}"
        _notify_status(status_label, text=_msg_ok, text_color="green")
        log_data(f"整理完成，成功匹配{success_cnt}，未匹配{fail_cnt}条，输出={stem}")
        if show_info:
            _info_ok = f"已保存：{out_xlsx.name} / {out_csv.name}\n模糊阈值 {ft}，成功匹配 {success_cnt}，未匹配 {fail_cnt}"
            show_info("完成", _info_ok)
        return str(out_xlsx)
    except Exception as e:
        _err = str(e)
        _notify_status(status_label, text=f"❌ 整理失败：{_err}", text_color="red")
        log_data(f"[整理失败] {e}")
        if show_error:
            show_error("整理异常", f"发生异常：{_err}")


def generate_project_stat_tables(
    input_file,
    output_dir,
    status_label=None,
    *,
    output_filename="示例项目数据统计表.xlsx",
    col_width=20,
    row_height=25,
    header_height=28,
    activity_types=None,
    header_fill_hex="C5E0B4",
    sum_fill_hex="FFF2CC",
    total_fill_hex="FEDB61",
    font_header_name="微软雅黑",
    font_header_size=12,
    font_body_name="微软雅黑",
    font_body_size=11,
):
    try:
        if activity_types:
            activity_types = [str(x).strip() for x in activity_types if str(x).strip()]
        if not activity_types:
            activity_types = list(_DEFAULT_STAT_ACTIVITY_TYPES)
        all_types = list(dict.fromkeys(activity_types + ["未分类"]))
        ROW_HEIGHT = _stat_safe_int(row_height, 25)
        HEADER_HEIGHT = _stat_safe_int(header_height, 28)
        hf = _stat_sanitize_hex(header_fill_hex, "C5E0B4")
        sf = _stat_sanitize_hex(sum_fill_hex, "FFF2CC")
        tf = _stat_sanitize_hex(total_fill_hex, "FEDB61")
        header_fill = PatternFill("solid", fgColor=hf)
        sum_fill = PatternFill("solid", fgColor=sf)
        total_fill = PatternFill("solid", fgColor=tf)
        fhs = _stat_safe_int(font_header_size, 12)
        fbs = _stat_safe_int(font_body_size, 11)
        font_header = Font(name=(font_header_name or "微软雅黑")[:31], bold=True, size=fhs)
        font_body = Font(name=(font_body_name or "微软雅黑")[:31], size=fbs)
        col_w = _stat_safe_float(col_width, 20)
        if col_w <= 0:
            col_w = 20
        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        border = Border(left=Side("thin"), right=Side("thin"), top=Side("thin"), bottom=Side("thin"))

        df = pd.read_excel(input_file)
        for col in ["学校名称", "活动类型"]:
            if col not in df.columns:
                raise ValueError(f"数据表缺少字段：{col}")

        def norm_val(x):
            return str(x).replace("\u3000", "").strip() if pd.notna(x) else ""

        df["学校名称"] = df["学校名称"].apply(norm_val)
        df["活动类型"] = df["活动类型"].apply(normalize_activity_types).apply(norm_val)
        df_detail = df[df["学校名称"] != ""].copy()
        if df_detail.empty:
            raise ValueError("没有可统计的数据，请填写至少一条含学校名称的活动记录。")

        def split_items(text):
            if pd.isnull(text) or str(text).strip() in ["", "nan", "None"]:
                return []
            return list(dict.fromkeys(t.strip() for t in re.split(r"[，,、/|;；\s]", str(text)) if t.strip()))

        table1 = df_detail.copy()
        if "序号" not in table1.columns:
            table1.insert(0, "序号", range(1, len(table1) + 1))

        table2 = df_detail.groupby("学校名称").size().reset_index(name="活动总数")
        table2 = table2.sort_values("活动总数", ascending=False).reset_index(drop=True)
        table2.insert(0, "序号", range(1, len(table2) + 1))
        table2 = pd.concat(
            [table2, pd.DataFrame([{"序号": "", "学校名称": "合计", "活动总数": table2["活动总数"].sum()}])],
            ignore_index=True,
        )

        # 与「工具集V4.0_合并经典版」generate_project_stat_tables 表3/表4/表5 逻辑一致（iterrows）
        rows3, type_err_list, err_type_cnt = [], [], 0
        for _, row in df_detail.iterrows():
            ts = split_items(row["活动类型"])
            if not ts:
                rows3.append({"学校名称": row["学校名称"], "活动类型": "未分类"})
            else:
                assigned = set()
                for t in ts:
                    category = t if t in all_types else '未分类'
                    if category not in assigned:
                        rows3.append({'学校名称': row['学校名称'], '活动类型': category})
                        assigned.add(category)
                    if t not in all_types:
                        err_type_cnt += 1
                        type_err_list.append(t)
        df_exp3 = pd.DataFrame(rows3)
        table3 = pd.pivot_table(df_exp3, index="学校名称", columns="活动类型", aggfunc="size", fill_value=0)
        table3 = table3.reindex(columns=[x for x in all_types if x in table3.columns], fill_value=0)
        table3["合计"] = table3.sum(axis=1)
        table3 = table3.reset_index()
        total_row3 = table3.iloc[:, 1:].sum().to_frame().T
        total_row3["学校名称"] = "总计"
        table3_body = table3[table3["学校名称"] != "总计"].sort_values("合计", ascending=False).reset_index(drop=True)
        table3_body.insert(0, "序号", range(1, len(table3_body) + 1))
        table3 = pd.concat([table3_body, total_row3], ignore_index=True)

        table4, table5 = pd.DataFrame(), pd.DataFrame()
        if all(c in df.columns for c in ["覆盖专业", "服务学生（人数）", "服务教师（人数）"]):
            rows4, rows5 = [], []
            for _, row in df_detail.iterrows():
                majors = split_items(row.get("覆盖专业", ""))
                try:
                    s_total = parse_headcount(row.get('服务学生（人数）'))
                    t_total = parse_headcount(row.get('服务教师（人数）'))
                except ValueError as exc:
                    raise ValueError(f'第 {row.name + 2} 行：{exc}') from exc
                m_list = majors if majors else ["未填/全院"]
                num_m = len(m_list)
                for m in m_list:
                    rows4.append({"专业名称": m, "覆盖场次": 1, "累计学生人次(全额)": s_total, "累计教师人次(全额)": t_total})
                    rows5.append({"专业名称": m, "覆盖场次": 1, "分配学生人数(均分)": s_total / num_m, "分配教师人数(均分)": t_total / num_m})
            for t_df, r_list in [(table4, rows4), (table5, rows5)]:
                res = pd.DataFrame(r_list).groupby("专业名称").sum().reset_index()
                res = res.sort_values("覆盖场次", ascending=False).reset_index(drop=True)
                res.insert(0, "序号", range(1, len(res) + 1))
                sum_row = pd.DataFrame([{"序号": "", "专业名称": "合计", "覆盖场次": res["覆盖场次"].sum(), res.columns[3]: res.iloc[:, 3].sum(), res.columns[4]: res.iloc[:, 4].sum()}])
                if t_df is table4:
                    table4 = pd.concat([res, sum_row], ignore_index=True)
                else:
                    table5 = pd.concat([res, sum_row], ignore_index=True)

        if not os.path.exists(output_dir):
            os.makedirs(output_dir, exist_ok=True)
        _fn = (output_filename or "示例项目数据统计表.xlsx").strip() or "示例项目数据统计表.xlsx"
        _fn = os.path.basename(_fn.replace("\\", "/").split("/")[-1])
        for ch in r'\/:*?"<>|':
            _fn = _fn.replace(ch, "")
        if not _fn.lower().endswith(".xlsx"):
            _fn += ".xlsx"
        output_file = new_output_path(output_dir, _fn[:-5][:115] + '.xlsx', (input_file,))
        with staged_outputs(output_file) as (staged_file,):
            with pd.ExcelWriter(staged_file, engine="openpyxl") as writer:
                table1.to_excel(writer, index=False, sheet_name="表1_活动明细")
                table2.to_excel(writer, index=False, sheet_name="表2_学校活动总数")
                table3.to_excel(writer, index=False, sheet_name="表3_类型分布统计")
                if not table4.empty:
                    table4.to_excel(writer, index=False, sheet_name="表4_专业覆盖(全额)")
                    table5.to_excel(writer, index=False, sheet_name="表5_专业覆盖(均分)")

            # 与经典版一致：逐行设置行高（避免部分环境下 sheet_format.defaultRowHeight 表现不一致）
            wb = openpyxl.load_workbook(staged_file)
            for sn in wb.sheetnames:
                ws = wb[sn]
                ncol = ws.max_column
                nrow = ws.max_row
                ws.row_dimensions[1].height = HEADER_HEIGHT
                for cell in ws[1]:
                    cell.font, cell.alignment, cell.fill, cell.border = font_header, align_center, header_fill, border
                for r_idx, row in enumerate(ws.iter_rows(min_row=2), 2):
                    ws.row_dimensions[r_idx].height = ROW_HEIGHT
                    for cell in row:
                        cell.font, cell.alignment, cell.border = font_body, align_center, border
                        if r_idx == nrow:
                            cell_val = str(ws.cell(nrow, 2).value)
                            cell.fill = total_fill if ("总计" in cell_val or "合计" in cell_val) else sum_fill
                        if sn == "表3_类型分布统计" and cell.column == ncol:
                            cell.fill = sum_fill
                for col in ws.columns:
                    ws.column_dimensions[get_column_letter(col[0].column)].width = col_w
            wb.save(staged_file)

        msg = f"✅ 统计表已输出（异常类型{err_type_cnt}条已计入未分类）"
        if err_type_cnt > 0:
            msg += f"\n异常示例: {set(type_err_list[:3])}"
        _notify_status(status_label, text=msg, text_color="green")
        if show_info:
            _ec = err_type_cnt
            generated = '已生成 5 张表，含专业覆盖全额及均分表。' if not table4.empty else '已生成 3 张表；缺少覆盖专业或师生人数列，未生成专业覆盖表。'
            _si = f"统计表生成完毕！\n{generated}\n共发现{_ec}条异常标签\n保存为：{Path(output_file).name}"
            show_info("完成", _si)
        return str(output_file)
    except Exception as e:
        _err = str(e)
        _notify_status(status_label, text=f"❌ 统计失败：{_err}", text_color="red")
        if show_error:
            show_error("统计失败", f"发生异常：{_err}")
