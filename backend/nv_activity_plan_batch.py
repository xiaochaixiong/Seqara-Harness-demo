# -*- coding: utf-8 -*-
"""自「发票分类」导出工作簿中的封面批量生成活动方案：产教融合分组封面逐张生成；机票报销封面整本合并为 1 份。"""
from __future__ import annotations

import re
import shutil
import threading
from difflib import SequenceMatcher
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, List, Optional

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

import nv_activity_plan_core as nv_plan
import nv_activity_plan_docx as nv_ap_docx
import nv_cover_fill_core as nv_cf

COVER_TITLE_TOKEN = "产教融合项目"
DETAIL_SUFFIXES = list(nv_cf.SETTINGS.get("detail_name_suffixes", ["明细", "-明细"]))


_MONEY_PAT = re.compile(
    r"\b\d{1,3}(?:,\d{3})*(?:\.\d+)?\s*(?:元|RMB|人民币|万元|千元|亿元)\b"
    r"|(\b\d{1,3}(?:,\d{3})*(?:\.\d+)?\b\s*（?:元|RMB|人民币|万元|千元|亿元）\b)"
    r"|\b\d+(?:\.\d+)?\s*(?:元|RMB|人民币)\b",
    re.IGNORECASE,
)
_INVOICE_NO_PAT = re.compile(r"(\b\d{6,}\b)|(发票号码|发票号|发票)\s*[:：]?\s*\S+", re.IGNORECASE)


def _strip_money_and_invoice(raw: str) -> str:
    if raw is None:
        return ""
    s = str(raw)
    s = _MONEY_PAT.sub("（金额另行核定）", s)
    s = _INVOICE_NO_PAT.sub("", s)
    return s.strip()


def _is_valid_activity_time(s: str) -> bool:
    s = (s or "").strip()
    # 允许：2026年2月 / 2026年1月~2026年3月
    if not s:
        return False
    return bool(
        re.match(r"^\d{4}年\d{1,2}月(?:~\d{4}年\d{1,2}月)?$", s)
    )


def _sanitize_prompt_fields(
    *,
    act_type: str,
    college: str,
    act_time: str,
) -> tuple[str, str, str]:
    """
    加固：确保送入DeepSeek prompt 的字段不会携带金额/发票号。
    同时对 act_time 做格式校验，防止把脏数据带进 prompt。
    """
    act_type_s = _strip_money_and_invoice(act_type)
    college_s = _strip_money_and_invoice(college)
    act_time_s = _strip_money_and_invoice(act_time)

    act_type_s = re.sub(r"\s+", " ", act_type_s).strip()
    college_s = re.sub(r"\s+", " ", college_s).strip()
    act_time_s = re.sub(r"\s+", "", act_time_s).strip()

    # 时间格式不对：直接置空，后续会跳过该封面
    if not _is_valid_activity_time(act_time_s):
        act_time_s = ""

    # 活动类型尽量回落到受控类别（避免 C4 被误填成包含金额的长串）
    all_cats = list(nv_cf.PROJECT_CATEGORIES.get("all") or [])
    # 如果 act_type 本身是多个类别，用连接符拆一下再过滤
    # 允许：逗号/顿号/、/空格分隔
    tokens = [t.strip() for t in re.split(r"[、,，\s]+", act_type_s) if t.strip()]
    kept = [t for t in tokens if t in all_cats]
    if kept:
        act_type_s = nv_cf.SETTINGS.get("schools_joiner", "、").join(dict.fromkeys(kept))
    return act_type_s, college_s, act_time_s


def _is_detail_sheet(name: str) -> bool:
    return nv_cf.is_detail_sheet(name, list(DETAIL_SUFFIXES))


def _is_cover_sheet(name: str) -> bool:
    if COVER_TITLE_TOKEN not in name:
        return False
    return not _is_detail_sheet(name)


def _strip_ai_disclaimer_lines(text: str) -> str:
    if not text:
        return text
    lines = []
    for line in text.splitlines():
        s = line.strip()
        if re.search(r"文档部分内容可能由\s*AI\s*生成", s, re.I):
            continue
        if re.search(r"由\s*AI\s*辅助生成", s, re.I) and "注" in s:
            continue
        lines.append(line)
    return "\n".join(lines).strip()


def project_title_for_cover_e4(body: str) -> str:
    """封面「项目明细」格：取项目名称/活动名称，去掉文末「策划方案」等后缀。"""
    t = nv_ap_docx.extract_project_title_from_body(body)
    if not t:
        for line in (body or "").splitlines():
            s = line.strip().replace("\r", "")
            if not s:
                continue
            s = re.sub(r"^#+\s*", "", s).strip()
            # 去掉可能由模型/渲染引入的序号前缀（如“一、”）
            s = re.sub(r"^\s*[（(]?[一二三四五六七八九十0-9]+[)）.、]\s*", "", s).strip()
            if s.startswith("---"):
                continue
            m = re.search(r"(项目名称|活动名称)\s*[：:]\s*(.+)$", s)
            if m:
                s = m.group(2).strip()
            if len(s) >= 4:
                t = s
                break
    t = (t or "").strip()
    t = re.sub(r"\s*[。；;]\s*$", "", t)
    t = re.sub(r"(活动策划方案|活动方案|策划方案|项目方案)\s*$", "", t).strip()
    return t


def _normalize_title_for_batch_compare(s: str) -> str:
    t = (s or "").strip()
    t = re.sub(r"\s+", "", t)
    return t


def _batch_title_too_similar(candidate: str, previous: List[str], *, ratio: float = 0.85) -> bool:
    """与本批已生成标题比对：全等或序列相似度 ≥ ratio 视为过近。"""
    c = _normalize_title_for_batch_compare(candidate)
    if len(c) < 4:
        return False
    for p in previous:
        pn = _normalize_title_for_batch_compare(p)
        if len(pn) < 4:
            continue
        if c == pn:
            return True
        if SequenceMatcher(None, c, pn).ratio() >= ratio:
            return True
    return False


def _find_detail_sheet_name_for_cover(cover_name: str, wb) -> str | None:
    """
    根据封面命名规则反推对应明细 sheet 名称。
    nv_classify_core 生成逻辑：detail_name = f"{cover_name}明细"
    （部分历史数据可能出现 "-明细" 形态，因此做兜底候选）。
    """
    # 优先精确拼接
    for suf in ("明细", "-明细"):
        cand = f"{cover_name}{suf}"
        if cand in wb.sheetnames and _is_detail_sheet(cand):
            return cand

    # 兜底：同前缀且为明细后缀的任意 sheet
    for n in wb.sheetnames:
        if n != cover_name and n.startswith(cover_name) and _is_detail_sheet(n):
            return n
    return None


def read_cover_fields(ws: Worksheet) -> dict[str, str]:
    """读取分组封面 C4~F4（与 nv_classify / nv_cover_fill 模板一致）。"""
    def cell_str(ref: str) -> str:
        v = nv_cf.merged_top_left_cell(ws, ref).value
        if v is None:
            return ""
        return str(v).strip()

    return {
        "项目类别": cell_str("C4"),
        "服务院校": cell_str("D4"),
        "项目明细": cell_str("E4"),
        "时间": cell_str("F4"),
    }


_TIME_MONTH_SEG_PAT = re.compile(r"(\d{4})年(\d{1,2})月")


def _year_month_tuples_from_time_text(s: str) -> list[tuple[int, int]]:
    return [(int(y), int(m)) for y, m in _TIME_MONTH_SEG_PAT.findall(s or "")]


def _aggregate_ticket_cover_time_range(wb, cover_sheet_names: list[str]) -> str:
    """汇总各机票封面 F4「时间」中的年月（支持单月至区间），得到整本最小月～最大月。"""
    acc: list[tuple[int, int]] = []
    for name in cover_sheet_names:
        f = read_cover_fields(wb[name])
        acc.extend(_year_month_tuples_from_time_text(f["时间"]))
    if not acc:
        return ""
    mn, mx = min(acc), max(acc)
    if mn == mx:
        return f"{mn[0]}年{mn[1]}月"
    return f"{mn[0]}年{mn[1]}月~{mx[0]}年{mx[1]}月"


def _collect_buyers_from_ticket_details(
    wb, ticket_pairs: list[tuple[str, str, int]]
) -> list[str]:
    buyers: list[str] = []
    for _cover, detail_name, _idx in ticket_pairs:
        dws = wb[detail_name]
        header = nv_cf.find_ticket_detail_header_row(dws)
        if not header:
            continue
        hr = int(header["_row"])
        buyer_c = int(header["_buyer_col"])
        for r in range(hr + 1, dws.max_row + 1):
            bv = dws.cell(row=r, column=buyer_c).value
            if bv is not None and str(bv).strip():
                buyers.append(str(bv).strip())
    return buyers


def _all_school_names_for_ticket_company(
    canon: str,
    schools_map: dict[str, list[str]],
    joiner: str,
) -> str:
    pool = schools_map.get(canon) or []
    out: list[str] = []
    seen: set[str] = set()
    for sch in pool:
        sch = (sch or "").strip()
        if sch and sch not in seen:
            seen.add(sch)
            out.append(sch)
    return joiner.join(out)


def _sanitize_ticket_merged_prompt_fields(
    act_type: str,
    college: str,
    act_time: str,
) -> tuple[str, str, str]:
    """机票整本汇总：不按产教类别表过滤「项目类别」，仅剥离金额/发票号并校验活动时间格式。"""
    act_type_s = _strip_money_and_invoice(act_type)
    college_s = _strip_money_and_invoice(college)
    act_time_s = _strip_money_and_invoice(act_time)
    act_type_s = re.sub(r"\s+", " ", act_type_s).strip()
    college_s = re.sub(r"\s+", " ", college_s).strip()
    act_time_s = re.sub(r"\s+", "", act_time_s).strip()
    if not _is_valid_activity_time(act_time_s):
        act_time_s = ""
    return act_type_s, college_s, act_time_s


def run_batch_from_invoice_workbook(
    input_xlsx: Path,
    output_dir: Path,
    *,
    major_name: str,
    style_mode: str = "豆包风",
    model: Optional[str] = None,
    strip_detail_sheets: bool = False,
    log_cb: Optional[Callable[[str], None]] = None,
    cancel_event: Optional[threading.Event] = None,
) -> dict[str, Any]:
    """
    从 input_xlsx 识别产教融合「封面」sheet，逐张按 C4/D4/F4 调用DeepSeek生成正文并写回 E4。

    若存在成对的「机票报销封面」与「机票报销明细」，则在同一工作簿末尾再生成 **1 份** 汇总方案：
    项目类别与机票封面 C4 一致（四项固定）；服务院校为该公司在设置中的 **全部** 院校（不按城市筛选）；
    时间段为**工作簿内识别到的全部机票封面** F4 年月之最小～最大范围。
    写回：同一「项目明细/方案名称」写入**每一个**机票封面 E4（不仅限于有成对明细的序号）。

    不向 API 发送 Excel 或明细表；仅发送与单次生成相同的文本提示词。
    """
    def log(msg: str) -> None:
        if log_cb:
            log_cb(msg)

    input_xlsx = Path(input_xlsx)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    major_name = (major_name or "").strip()
    if not major_name:
        raise ValueError("批量模式需要填写「专业名称」（与单次生成一致）。")

    wb_probe = load_workbook(input_xlsx, read_only=True, data_only=True)
    try:
        probe_names = list(wb_probe.sheetnames)
    finally:
        wb_probe.close()

    ticket_pairs = nv_cf.iter_ticket_cover_detail_pairs(probe_names)
    _ticket_cover_tabs = nv_cf.list_ticket_cover_sheet_names_ordered(probe_names)
    cover_names = [n for n in probe_names if _is_cover_sheet(n)]
    ticket_job = 1 if ticket_pairs else 0
    if not cover_names and not ticket_pairs:
        raise RuntimeError(
            "未找到可批量生成的封面：既无名称含「产教融合项目」的分组封面，"
            "也无成对的「机票报销封面」与「机票报销明细」sheet。"
            "请确认已使用发票分类整理导出。"
        )

    if _ticket_cover_tabs and not ticket_pairs:
        log(
            "⚠️ 检测到机票类封面 sheet（名称形如「机票报销封面_n」或「…报销明细-1-1」等）"
            f"共 {len(_ticket_cover_tabs)} 个，但与「机票报销明细_n」未成对，已跳过机票汇总方案（不会写 E4）。"
            "请确认工作簿内仍有「机票报销明细_1」等明细表且未被改名。"
        )

    out_wb_path = output_dir / f"{input_xlsx.stem}_批量方案回填.xlsx"
    shutil.copy2(input_xlsx, out_wb_path)
    wb = load_workbook(out_wb_path)

    style_mode = (style_mode or "豆包风").strip() or "豆包风"
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")

    ok_count = 0
    skip_count = 0
    errors: List[str] = []
    recent_titles: List[str] = []

    for idx, cover_name in enumerate(cover_names, start=1):
        if cancel_event is not None and cancel_event.is_set():
            log(">>> 已取消，停止后续封面。")
            break
        ws = wb[cover_name]
        fields = read_cover_fields(ws)
        act_type = fields["项目类别"]
        college = fields["服务院校"]
        act_time = fields["时间"]

        # 加固：保证送入 API 的字段不会包含金额/发票号，并校验时间格式
        act_type, college, act_time = _sanitize_prompt_fields(
            act_type=act_type, college=college, act_time=act_time
        )

        log(f"\n>>> [{idx}/{len(cover_names)}] 封面：{cover_name}")
        if not act_type or not college or not act_time:
            msg = "跳过：封面 C4（项目类别）/D4（服务院校）/F4（时间）需已由「同步填写封面汇总」或手工填好。"
            log(f"⚠️ {msg}")
            skip_count += 1
            errors.append(f"{cover_name}: {msg}")
            continue
        try:
            var_key = f"{idx:04d}|{cover_name}"
            avoid = list(recent_titles) if recent_titles else None
            body = nv_plan.generate_activity_plan_body(
                act_type,
                major_name,
                act_time,
                college,
                style_mode=style_mode,
                log_cb=log_cb,
                model=model,
                variation_key=var_key,
                batch_slot=idx,
                batch_total=len(cover_names),
                avoid_titles=avoid,
            )
            body = _strip_ai_disclaimer_lines(body)
            e4_text = project_title_for_cover_e4(body)
            if not e4_text:
                e4_text = nv_ap_docx.extract_project_title_from_body(body) or "（请从 Word 正文核对主题）"

            _ph = "（请从 Word 正文核对主题）"
            first_e4 = e4_text
            if first_e4 and first_e4 != _ph and _batch_title_too_similar(first_e4, recent_titles):
                log("⚠️ 项目名称与同批已有方案过于接近，自动重试 1 次…")
                avoid_retry = list(recent_titles)
                avoid_retry.append(first_e4)
                body = nv_plan.generate_activity_plan_body(
                    act_type,
                    major_name,
                    act_time,
                    college,
                    style_mode=style_mode,
                    log_cb=log_cb,
                    model=model,
                    variation_key=f"{var_key}|dupretry1",
                    batch_slot=idx,
                    batch_total=len(cover_names),
                    avoid_titles=avoid_retry,
                    extra_user_suffix=nv_plan.ACTIVITY_PLAN_TITLE_DEDUP_RETRY_SUFFIX,
                )
                body = _strip_ai_disclaimer_lines(body)
                e4_text = project_title_for_cover_e4(body)
                if not e4_text:
                    e4_text = nv_ap_docx.extract_project_title_from_body(body) or _ph
                if e4_text and e4_text != _ph:
                    if _batch_title_too_similar(e4_text, recent_titles) or _batch_title_too_similar(
                        e4_text, [first_e4]
                    ):
                        log("⚠️ 重试后标题仍与某条同批标题或首次标题较接近，已保留重试结果。")

            nv_cf.set_cover_block_style(ws)
            nv_cf.write_merged_top_left(ws, "E4", e4_text)

            # 同步回填到“原始包含明细的表格”（明细 sheet 标题 B1）
            detail_name = _find_detail_sheet_name_for_cover(cover_name, wb)
            if detail_name:
                try:
                    ws_detail = wb[detail_name]
                    old = ws_detail["B1"].value
                    old_s = str(old or "")
                    if e4_text and e4_text not in old_s:
                        ws_detail["B1"].value = f"{old_s}\n{e4_text}".strip()
                        log(f"✅ 已写回明细标题（{detail_name}）：{e4_text[:40]}{'…' if len(e4_text) > 40 else ''}")
                except Exception as e:
                    log(f"⚠️ 明细标题写回失败（{detail_name}）：{e}")

            doc_stem = nv_ap_docx.safe_filename(f"{idx:02d}_{cover_name}_{ts}")
            doc_path = output_dir / f"{doc_stem}.docx"
            nv_ap_docx.save_activity_plan_to_docx(
                str(doc_path),
                "",
                college,
                act_type,
                act_time,
                "",
                body,
                cover_title_override=e4_text or None,
            )
            log(f"✅ 已生成 Word：{doc_path.name}")
            log(f"✅ 已写回封面 E4（项目明细）：{e4_text[:80]}{'…' if len(e4_text) > 80 else ''}")
            ok_count += 1
            if e4_text and e4_text != "（请从 Word 正文核对主题）":
                recent_titles.append(e4_text)
        except Exception as e:
            err = f"{cover_name}: {e}"
            errors.append(err)
            log(f"❌ {err}")

    if ticket_pairs and (cancel_event is None or not cancel_event.is_set()):
        log(
            "\n>>> 机票报销：整本合并为 1 份活动方案（服务院校=该公司全部院校；"
            "时间=各机票封面 F4 年月的最小～最大）…"
            f"（已配对明细 {len(ticket_pairs)} 组）"
        )
        # 与「成对明细」无关：凡识别到的机票封面均参与 F4 时间汇总，且写回同一套「项目明细」文案
        all_ticket_cover_names = nv_cf.list_ticket_cover_sheet_names_ordered(list(wb.sheetnames))
        cover_sheet_names = [c for c, _d, _i in ticket_pairs]
        covers_for_time = all_ticket_cover_names if all_ticket_cover_names else cover_sheet_names
        covers_for_e4 = all_ticket_cover_names if all_ticket_cover_names else cover_sheet_names
        buyers = _collect_buyers_from_ticket_details(wb, ticket_pairs)
        schools_map = nv_cf.load_ticket_company_schools_runtime()
        canon, cerr = nv_cf._canonical_ticket_company(buyers, None, schools_map)
        if cerr or not canon:
            msg = cerr or "无法确定机票购买方公司"
            log(f"⚠️ 跳过机票汇总方案：{msg}")
            skip_count += 1
            errors.append(f"机票汇总: {msg}")
        else:
            joiner = str(nv_cf.SETTINGS.get("schools_joiner") or "、")
            college_all = _all_school_names_for_ticket_company(canon, schools_map, joiner)
            act_type = joiner.join(nv_cf.TICKET_COVER_CATEGORY_LABELS)
            act_time = _aggregate_ticket_cover_time_range(wb, covers_for_time)
            act_type, college_all, act_time = _sanitize_ticket_merged_prompt_fields(
                act_type, college_all, act_time
            )
            if not college_all:
                msg = f"公司「{canon}」在 ticket 规则中无院校列表，无法生成汇总方案"
                log(f"⚠️ 跳过机票汇总方案：{msg}")
                skip_count += 1
                errors.append(f"机票汇总: {msg}")
            elif not act_time:
                msg = "各机票封面 F4「时间」中未解析到有效年月（请先运行「同步填写机票封面汇总」或手工填写）"
                log(f"⚠️ 跳过机票汇总方案：{msg}")
                skip_count += 1
                errors.append(f"机票汇总: {msg}")
            else:
                try:
                    body = nv_plan.generate_activity_plan_body(
                        act_type,
                        major_name,
                        act_time,
                        college_all,
                        style_mode=style_mode,
                        log_cb=log_cb,
                        model=model,
                        variation_key="机票报销汇总",
                    )
                    body = _strip_ai_disclaimer_lines(body)
                    e4_text = project_title_for_cover_e4(body)
                    if not e4_text:
                        e4_text = (
                            nv_ap_docx.extract_project_title_from_body(body)
                            or "（请从 Word 正文核对主题）"
                        )
                    for cname in covers_for_e4:
                        try:
                            cws = wb[cname]
                            nv_cf.set_cover_block_style(cws)
                            nv_cf.write_merged_top_left(cws, "E4", e4_text)
                        except Exception as e:
                            log(f"⚠️ 写回机票封面 E4 失败（{cname}）：{e}")
                    for _c, detail_name, _i in ticket_pairs:
                        try:
                            ws_detail = wb[detail_name]
                            old = ws_detail["B1"].value
                            old_s = str(old or "")
                            if e4_text and e4_text not in old_s:
                                ws_detail["B1"].value = f"{old_s}\n{e4_text}".strip()
                                log(
                                    f"✅ 已写回机票明细标题（{detail_name}）："
                                    f"{e4_text[:40]}{'…' if len(e4_text) > 40 else ''}"
                                )
                        except Exception as e:
                            log(f"⚠️ 机票明细标题写回失败（{detail_name}）：{e}")
                    batch_idx = len(cover_names) + 1
                    doc_stem = nv_ap_docx.safe_filename(
                        f"{batch_idx:02d}_机票报销汇总_{ts}"
                    )
                    doc_path = output_dir / f"{doc_stem}.docx"
                    nv_ap_docx.save_activity_plan_to_docx(
                        str(doc_path),
                        "",
                        college_all,
                        act_type,
                        act_time,
                        "",
                        body,
                        cover_title_override=e4_text or None,
                    )
                    log(f"✅ 已生成机票汇总 Word：{doc_path.name}")
                    log(
                        f"✅ 已写回 {len(covers_for_e4)} 个机票封面 E4（同一项目明细）；"
                        f"服务院校数：{len(college_all.split(joiner)) if college_all else 0}；时间：{act_time}"
                    )
                    ok_count += 1
                except Exception as e:
                    err = f"机票汇总: {e}"
                    errors.append(err)
                    log(f"❌ {err}")

    if strip_detail_sheets:
        # 不删除 sheet：避免你看到“封面后面没有明细”，同时通过隐藏明细行实现脱敏。
        to_strip = [n for n in list(wb.sheetnames) if _is_detail_sheet(n)]
        log(f"\n>>> 脱敏：保留 {len(to_strip)} 个明细 sheet，仅隐藏明细行…")
        for name in to_strip:
            try:
                ws_d = wb[name]
                # 明细数据从第 3 行开始（第 1 行：标题；第 2 行：表头）
                for r in range(3, (ws_d.max_row or 0) + 1):
                    ws_d.row_dimensions[r].hidden = True
            except Exception:
                # 不让脱敏步骤影响主流程
                continue

    wb.save(out_wb_path)

    return {
        "output_workbook": str(out_wb_path),
        "output_dir": str(output_dir),
        "covers_total": len(cover_names) + ticket_job,
        "ok": ok_count,
        "skipped": skip_count,
        "errors": errors,
    }
