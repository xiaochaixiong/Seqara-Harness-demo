import importlib.util
from pathlib import Path
import tempfile
import unittest
import zipfile

module_path=Path(__file__).resolve().parents[1]/'backend/presentation_io.py'
spec=importlib.util.spec_from_file_location('presentation_io',module_path)
io=importlib.util.module_from_spec(spec);spec.loader.exec_module(io)

class MaterialTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(prefix='seqara-material-');self.folder=Path(self.temp.name)
    def tearDown(self): self.temp.cleanup()
    def test_text_pagination_preserves_scope_and_hash(self):
        file=self.folder/'材料.md';file.write_text('\n'.join(f'测试项 {i}' for i in range(180)),encoding='utf-8')
        first=io.read_source({'source':str(file),'limit':100});second=io.read_source({'source':str(file),'offset':first['next_offset']})
        self.assertEqual(first['total_fragments'],180);self.assertEqual(len(second['fragments']),80)
        self.assertIsNone(second['next_offset']);self.assertEqual(first['source_hash'],second['source_hash'])
        self.assertEqual(second['fragments'][0]['location'],'行 101')
    def test_docx_and_excel_keep_text_location_and_formula(self):
        from docx import Document
        from openpyxl import Workbook
        doc=Document();doc.add_paragraph('团队交接测试');table=doc.add_table(rows=2,cols=2)
        table.cell(0,0).text='阶段';table.cell(0,1).text='分钟';table.cell(1,0).text='试行';table.cell(1,1).text='30'
        file=self.folder/'材料.docx';doc.save(file)
        fragments=io.read_source({'source':str(file)})['fragments'];self.assertIn('试行 | 30',str(fragments));self.assertIn('表格行 2',str(fragments))
        book=Workbook();sheet=book.active;sheet.title='测试';sheet.append(['时长',45,30,'=B1-C1']);file=self.folder/'材料.xlsx';book.save(file)
        result=io.read_source({'source':str(file)});self.assertIn('D1==B1-C1',result['fragments'][0]['text']);self.assertIn('测试!行1',result['fragments'][0]['location']);self.assertIn('没有',result['notes'])
    def test_unsupported_and_out_of_range_are_actionable(self):
        file=self.folder/'old.ppt';file.write_bytes(b'placeholder')
        with self.assertRaisesRegex(ValueError,'材料支持'):io.read_source({'source':str(file)})
        file=self.folder/'data.txt';file.write_text('hello',encoding='utf8')
        with self.assertRaisesRegex(ValueError,'offset'):io.read_source({'source':str(file),'offset':2})
    def test_pptx_text_scope_is_explicit(self):
        file=self.folder/'材料.pptx'
        with zipfile.ZipFile(file,'w') as archive:
            archive.writestr('ppt/slides/slide1.xml','<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:p><a:r><a:t>可见材料42</a:t></a:r></a:p></p:sld>')
        result=io.read_source({'source':str(file)})
        self.assertEqual(result['fragments'][0]['text'],'可见材料42')
        self.assertIn('备注',result['notes'])

if __name__=='__main__':unittest.main()
